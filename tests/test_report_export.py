from __future__ import annotations

from datetime import datetime

from openpyxl import load_workbook

from models.entities import UiKind
from models.report_export import (
    NO_RESPONSIBLE,
    SHEET_BY_RESP,
    SHEET_MAP,
    SHEET_RELATIONS,
    SHEET_SECTIONS,
    SHEET_SUMMARY,
    export_report_xlsx,
)


def _populate(model):
    sts = {s.name: s.id for s in model.statuses()}
    rs = {r.code: r.id for r in model.responsibles()}
    a = model.add_section("31 33 23", "Estabilización de roca")
    b = model.add_section("01 35 29", "Requisitos de seguridad")
    c = model.add_section("33 40 00", "Drenaje pluvial")
    d = model.add_section("03 30 00", "Concreto")
    e = model.add_section("4.28.33", "Sitio de obra")
    model.set_section_status(a.id, sts["En elaboración"]); model.set_section_progress(a.id, 70)
    model.set_section_responsibles(a.id, [rs["INIO"], rs["INIC"]]); model.set_section_observations(a.id, "Pendiente ensayo")
    model.set_section_status(c.id, sts["Aprobada"]); model.set_section_progress(c.id, 100)
    model.set_section_responsibles(c.id, [rs["INIC"]])
    model.set_section_progress(b.id, 40); model.set_section_responsibles(b.id, [rs["INIO"]])
    model.add_relation(a.id, UiKind.REFERENCES, b.id)
    model.add_relation(c.id, UiKind.REFERENCES, a.id)
    model.add_relation(d.id, UiKind.REFERENCES, e.id)
    model.add_relation(e.id, UiKind.REFERENCES, d.id)   # dos flechas opuestas
    return a, b, c, d, e


def test_report_has_expected_sheets_and_summary(tmp_path, model):
    _populate(model)
    out = tmp_path / "reporte.xlsx"
    stats = export_report_xlsx(model, out)
    assert out.exists() and stats.sections == 5 and stats.relations == 4
    assert stats.avg_progress == (70 + 40 + 100 + 0 + 0) / 5 and stats.completed == 1
    assert stats.without_responsible == 2
    wb = load_workbook(out)
    assert wb.sheetnames == [SHEET_SUMMARY, SHEET_SECTIONS, SHEET_BY_RESP, SHEET_RELATIONS]
    ws = wb[SHEET_SUMMARY]
    assert "CC-25-01" in ws["A1"].value and "Proyecto de prueba" in ws["A1"].value
    assert str(datetime.now().year) in ws["A2"].value
    labels = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(6, 14)}
    assert labels["Secciones"] == 5 and labels["Relaciones"] == 4
    assert labels["Avance promedio"] == 42 and labels["Secciones al 100 %"] == 1
    assert labels["Secciones sin responsable"] == 2
    texts = [str(ws.cell(row=r, column=1).value) for r in range(1, ws.max_row + 1)]
    assert "INIO" in texts and "INIC" in texts and NO_RESPONSIBLE in texts
    row_inic = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "INIC")
    assert ws.cell(row=row_inic, column=3).value == 2       # secciones de INIC
    assert ws.cell(row=row_inic, column=4).value == 85      # avance promedio (70 + 100) / 2
    assert ws.cell(row=row_inic, column=5).value == 1       # al 100 %


def test_sections_sheet_rows_and_formatting(tmp_path, model):
    a, *_ = _populate(model)
    out = tmp_path / "reporte.xlsx"
    export_report_xlsx(model, out)
    ws = load_workbook(out)[SHEET_SECTIONS]
    assert ws.max_row == 6 and ws["A1"].value == "N.º"
    rows = {ws.cell(row=r, column=2).value: r for r in range(2, 7)}
    r = rows["31 33 23"]
    assert ws.cell(row=r, column=5).value == "En elaboración"
    assert ws.cell(row=r, column=6).value == 70
    assert ws.cell(row=r, column=7).value == "INIO, INIC"
    assert ws.cell(row=r, column=8).value == "Pendiente ensayo"
    assert ws.cell(row=r, column=9).value == 1 and ws.cell(row=r, column=10).value == 1
    status = model.status(model.section(a.id).status_id)
    assert ws.cell(row=r, column=5).fill.fgColor.rgb.endswith(status.color.lstrip("#").upper())
    assert ws.auto_filter.ref and ws.freeze_panes == "C2"
    rules = [rng for rng in ws.conditional_formatting]
    assert any("F2" in str(rng.sqref) for rng in rules)


def test_by_responsible_and_relations_sheets(tmp_path, model):
    _populate(model)
    out = tmp_path / "reporte.xlsx"
    export_report_xlsx(model, out)
    wb = load_workbook(out)
    ws = wb[SHEET_BY_RESP]
    first_col = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert first_col.count("INIO") == 2 and first_col.count("INIC") == 2
    assert first_col.count(NO_RESPONSIBLE) == 2
    assert ws.max_row == 1 + 4 + 2
    rel = wb[SHEET_RELATIONS]
    arrows = [rel.cell(row=r, column=3).value for r in range(2, rel.max_row + 1)]
    assert arrows.count("→") == 4 and "↔" not in arrows


def test_map_sheet_when_image_given(tmp_path, model):
    from PIL import Image

    _populate(model)
    img = tmp_path / "mapa.png"
    Image.new("RGB", (200, 100), "white").save(img)
    out = tmp_path / "reporte.xlsx"
    stats = export_report_xlsx(model, out, img)
    assert SHEET_MAP in stats.sheets
    ws = load_workbook(out)[SHEET_MAP]
    assert "Mapa de referencias" in ws["A1"].value and len(ws._images) == 1


def test_empty_project_report(tmp_path, model):
    out = tmp_path / "vacio.xlsx"
    stats = export_report_xlsx(model, out)
    assert stats.sections == 0 and stats.avg_progress == 0
    wb = load_workbook(out)
    assert wb[SHEET_SECTIONS].max_row == 1


def test_clauses_are_excluded_from_progress_and_responsible_sheets(tmp_path, model):
    from openpyxl import load_workbook

    a = model.add_section("03 30 00", "Concreto")
    model.set_section_progress(a.id, 40)
    c = model.add_section("4.28.61", "PAGO FINAL", kind="clause")
    model.add_relation(a.id, UiKind.REFERENCES, c.id)
    out = tmp_path / "r.xlsx"
    stats = export_report_xlsx(model, out)
    assert stats.sections == 2 and stats.avg_progress == 40.0 and stats.without_responsible == 1
    wb = load_workbook(out)
    ws = wb[SHEET_SECTIONS]
    rows = {ws.cell(row=r, column=2).value: [ws.cell(row=r, column=col).value for col in range(1, 8)]
            for r in range(2, ws.max_row + 1)}
    assert rows["4.28.61"][3] == "Cláusula" and rows["4.28.61"][4] in ("", None) and rows["4.28.61"][5] is None
    assert rows["03 30 00"][5] == 40
    by_resp = wb[SHEET_BY_RESP]
    codes = [by_resp.cell(row=r, column=3).value for r in range(2, by_resp.max_row + 1)]
    assert "03 30 00" in codes and "4.28.61" not in codes
