"""Pruebas de humo de la interfaz completa (offscreen, pytest-qt)."""
from __future__ import annotations

import os

import pytest
from PyQt6.QtCore import QPointF, Qt

from controllers.export_controller import render_png, render_svg
from controllers.main_controller import MainController
from models.entities import RelationKind, UiKind
from models.masterformat_tree_model import ROLE_IN_PROJECT as ROLE_IN_PROJECT_TREE
from models.masterformat_tree_model import MasterFormatTreeModel
from models.project_model import ProjectModel
from models.relations_table_model import COL_KIND, RelationsTableModel
from models.section_completer_model import SectionCompleterModel
from views.main_window import MainWindow


@pytest.fixture
def app(qtbot, monkeypatch, master):
    # QMessageBox modales: responder "Sí" automáticamente.
    from PyQt6.QtWidgets import QMessageBox

    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.StandardButton.Yes))
    monkeypatch.setattr(QMessageBox, "information", staticmethod(lambda *a, **k: None))
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: None))

    project = ProjectModel(master=master)
    table_model = RelationsTableModel(project)
    completer_model = SectionCompleterModel()
    window = MainWindow(table_model, completer_model, MasterFormatTreeModel(master))
    controller = MainController(project, window, table_model, completer_model)
    qtbot.addWidget(window)
    project.new_project(None, "CC-25-01", "Demo")
    window.show()
    yield project, window, controller, table_model
    project.close()


def _add(window, a: str, kind: UiKind, b: str) -> None:
    """Simula la entrada en modo estricto: las secciones se eligen (se crean si hace falta)."""
    from models.relation_normalizer import split_code_title

    project = window.table_view.table.model().project
    for which, text in (("a", a), ("b", b)):
        code, title = split_code_title(text)
        section = project.section_by_code(code) or project.add_section(code, title)
        window.entry.picker(which).set_section(section.id, section.label)
    window.entry.kind_combo.setCurrentText(kind.value)
    window.entry.add_button.click()


def test_entry_creates_sections_row_and_canvas_items(app, qtbot):
    project, window, _c, table_model = app
    _add(window, "31 33 23 Estabilización de roca", UiKind.REFERENCES, "01 31 19 Conferencia inicial")
    assert table_model.rowCount() == 1
    assert len(window.scene.nodes) == 2
    assert len(window.scene.edges) == 1
    rel = project.relations()[0]
    sa, sb = project.section(rel.source_id), project.section(rel.target_id)
    assert sa.code == "31 33 23" and sb.code == "01 31 19"
    edge = window.scene.edges[rel.id]
    assert edge.effective_ports[0] in ("top", "right") and edge.effective_ports[1] in ("bottom", "left")


def test_referenced_by_is_stored_canonically(app):
    project, window, _c, _t = app
    _add(window, "33 40 00 Drenaje pluvial", UiKind.REFERENCED_BY, "31 23 00 Excavación")
    rel = project.relations()[0]
    assert project.section(rel.source_id).code == "31 23 00"
    assert project.section(rel.target_id).code == "33 40 00"
    assert rel.kind is RelationKind.REF


def test_duplicate_inverse_becomes_mutual(app):
    project, window, _c, table_model = app
    _add(window, "A", UiKind.REFERENCES, "B")
    _add(window, "B", UiKind.REFERENCES, "A")  # QMessageBox.question -> Yes (convertir en mutua)
    assert table_model.rowCount() == 1
    assert project.relations()[0].kind is RelationKind.MUTUAL
    edge = next(iter(window.scene.edges.values()))
    assert len(edge._arrows) == 2


def test_move_node_persists_and_edges_follow(app, qtbot):
    project, window, _c, _t = app
    _add(window, "A", UiKind.REFERENCES, "B")
    a = project.section_by_code("A")
    node = window.scene.nodes[a.id]
    edge = next(iter(window.scene.edges.values()))
    before = [QPointF(p) for p in edge.points]
    node.setPos(node.pos() + QPointF(300, 200))
    window.scene.nodesMoved.emit({a.id: (node.pos().x(), node.pos().y())})
    pos = project.position(a.id)
    assert pos.pinned and abs(pos.x - node.pos().x()) < 0.01
    assert edge.points != before
    # Cambio programático desde el modelo -> escena (sin bucle)
    project.move_node(a.id, 10.0, 20.0)
    assert node.pos() == QPointF(10.0, 20.0)


def test_delete_relation_from_table_removes_edge(app):
    project, window, controller, table_model = app
    _add(window, "A", UiKind.REFERENCES, "B")
    rid = project.relations()[0].id
    controller.relations.delete_relation(rid)
    assert table_model.rowCount() == 0
    assert rid not in window.scene.edges
    assert len(window.scene.nodes) == 2  # las secciones permanecen


def test_change_kind_from_table_updates_edge(app):
    project, window, _c, table_model = app
    _add(window, "A", UiKind.REFERENCES, "B")
    idx = table_model.index(0, COL_KIND)
    table_model.setData(idx, UiKind.MUTUAL.value, Qt.ItemDataRole.EditRole)
    rel = project.relations()[0]
    assert rel.kind is RelationKind.MUTUAL
    assert window.scene.edges[rel.id].kind is RelationKind.MUTUAL


def test_remove_section_removes_node_and_edges(app):
    project, window, controller, table_model = app
    _add(window, "A", UiKind.REFERENCES, "B")
    _add(window, "A", UiKind.REFERENCES, "C")
    a = project.section_by_code("A")
    controller.canvas.delete_section(a.id)
    assert a.id not in window.scene.nodes
    assert len(window.scene.edges) == 0
    assert table_model.rowCount() == 0


def test_waypoint_geometry_roundtrip(app):
    project, window, _c, _t = app
    _add(window, "A", UiKind.REFERENCES, "B")
    rel = project.relations()[0]
    edge = window.scene.edges[rel.id]
    mid = edge.points[len(edge.points) // 2]
    edge.insert_waypoint_at(mid + QPointF(0, 40))
    stored = project.relation(rel.id)
    assert stored.waypoints and len(stored.waypoints) == len(edge.waypoints)
    edge.reset_geometry()
    assert project.relation(rel.id).waypoints is None


def test_export_png_and_svg(app, tmp_path):
    _p, window, _c, _t = app
    _add(window, "A", UiKind.REFERENCES, "B")
    png = tmp_path / "mapa.png"
    svg = tmp_path / "mapa.svg"
    render_png(window.scene, png, scale=1.0)
    render_svg(window.scene, svg)
    assert png.stat().st_size > 0 and svg.stat().st_size > 0


def test_export_3d_png_and_svg(app, tmp_path, qtbot):
    from controllers.export_controller import render_3d_png, render_3d_svg
    from views.main_window import TAB_3D

    project, window, controller, _t = app
    _add(window, "A", UiKind.REFERENCES, "B")
    _add(window, "B", UiKind.REFERENCES, "C")
    window.tabs.setCurrentIndex(TAB_3D)
    if not window.view3d.available:
        pytest.skip("OpenGL no disponible en este entorno")
    controller.view3d.refresh()
    # La disposición se calcula en segundo plano la primera vez: esperar a que se renderice.
    qtbot.waitUntil(lambda: window.view3d._data is not None and len(window.view3d._data["ids"]) == 3,
                    timeout=60000)
    png, svg = tmp_path / "v3d.png", tmp_path / "v3d.svg"
    # El SVG es una proyección vectorial: no depende del framebuffer y debe funcionar siempre.
    render_3d_svg(window.view3d, svg)
    text = svg.read_text(encoding="utf-8")
    assert "<svg" in text and (text.count("<circle") >= 3 or "<path" in text)
    scene = window.view3d.projected_scene()
    assert len(scene["nodes"]) == 3 and len(scene["edges"]) == 2
    # El PNG requiere un framebuffer OpenGL real; en la plataforma offscreen no existe.
    if os.environ.get("QT_QPA_PLATFORM") == "offscreen":
        pytest.skip("Sin framebuffer OpenGL en plataforma offscreen")
    render_3d_png(window.view3d, png, scale=1.0)
    assert png.stat().st_size > 0


def test_projected_scene_shape_without_gl_data(app):
    """La proyección devuelve estructuras vacías coherentes cuando no hay datos."""
    from controllers.export_controller import paint_projected_scene
    from PyQt6.QtGui import QImage, QPainter

    scene = {"size": (200, 100), "nodes": [], "edges": []}
    image = QImage(200, 100, QImage.Format.Format_ARGB32)
    painter = QPainter(image)
    paint_projected_scene(painter, scene)
    painter.end()
    assert not image.isNull()


def test_analysis_panel_and_highlight(app):
    project, window, controller, _t = app
    _add(window, "A", UiKind.REFERENCES, "B")
    _add(window, "B", UiKind.REFERENCES, "C")
    project.add_section("Z", "Huérfana")
    controller.analysis.refresh()
    assert window.analysis.orphans_list.count() == 1
    assert window.analysis.hubs_table.rowCount() == 3
    c = project.section_by_code("C")
    controller.analysis.highlight(c.id)
    z = project.section_by_code("Z")
    assert window.scene.nodes[z.id].opacity() < 1.0
    assert window.scene.nodes[c.id].opacity() == 1.0
    controller.canvas.clear_highlight()
    assert window.scene.nodes[z.id].opacity() == 1.0


def test_view3d_layout_cached(app):
    project, window, controller, _t = app
    _add(window, "A", UiKind.REFERENCES, "B")
    positions = controller.view3d._positions()
    assert set(positions) == {s.id for s in project.sections()}
    again = controller.view3d._positions()
    assert positions == again
    _add(window, "B", UiKind.REFERENCES, "C")
    grown = controller.view3d._positions()
    assert len(grown) == 3 and all(grown[k] == positions[k] for k in positions)


def test_section_editor_dialog_builds_with_category_icons(app):
    from views.components.section_editor_dialog import SectionEditorDialog

    project, window, _c, _t = app
    dialog = SectionEditorDialog(project.categories(), window, is_new=True)
    assert dialog.category_combo.count() == len(project.categories()) + 1
    dialog.code_edit.setText("31 23 00")
    dialog.title_edit.setText("Excavación")
    assert dialog.values()[:2] == ("31 23 00", "Excavación")


def test_custom_section_color_reaches_canvas_and_table(app):
    from models.relations_table_model import COL_A, ROLE_FILL
    from views.components.section_editor_dialog import SectionEditorDialog

    project, window, _c, table_model = app
    _add(window, "A", UiKind.REFERENCES, "B")
    a = project.section_by_code("A")
    project.set_section_colors(a.id, "#FF8800")
    assert window.scene.nodes[a.id].fill.name().upper() == "#FF8800"
    assert table_model.data(table_model.index(0, COL_A), ROLE_FILL) == "#FF8800"
    dialog = SectionEditorDialog(project.categories(), window, code="A", title="", category_id=a.category_id,
                                 fill_color="#FF8800", border_color=project.section(a.id).border_color)
    assert dialog.custom_color_check.isChecked()
    assert dialog.colors()[0] == "#FF8800"
    dialog.custom_color_check.setChecked(False)
    assert dialog.colors() == (None, None)


def test_help_dialog_contextual_topic_and_first_time_hint(app):
    from views.main_window import TAB_3D, TAB_ANALYSIS, TAB_SECTIONS

    _p, window, controller, _t = app
    help_ctrl = controller.help
    assert help_ctrl.contextual_topic() == "mapa"
    window.tabs.setCurrentIndex(TAB_SECTIONS)
    assert help_ctrl.contextual_topic() == "estatus"
    window.tabs.setCurrentIndex(TAB_ANALYSIS)
    assert help_ctrl.contextual_topic() == "analisis"
    window.tabs.setCurrentIndex(TAB_3D)
    assert help_ctrl.contextual_topic() == "vista3d"
    help_ctrl.show_contextual()
    dialog = help_ctrl.dialog()
    assert dialog.isVisible() and dialog.current_topic == "vista3d"
    dialog.show_topic("relaciones")
    assert "Referencia mutua" in dialog.browser.toPlainText()
    dialog.search.setText("conectar")
    assert dialog.topics_list.count() >= 2
    dialog.close()
    # Pista de primera vez al activar Conectar (solo una vez por sesión)
    window.act_connect.setChecked(True)
    assert "Conectar" in window.statusBar().currentMessage()
    window.act_connect.setChecked(False)
    window.statusBar().clearMessage()
    window.act_connect.setChecked(True)
    assert window.statusBar().currentMessage() == ""
    window.act_connect.setChecked(False)


def test_about_dialog_shows_version_and_developer(app):
    from config.settings import APP_VERSION, DEVELOPER_NAME, GITHUB_USER
    from views.components.about_dialog import AboutDialog

    _p, window, _c, _t = app
    dialog = AboutDialog(window)
    assert APP_VERSION in dialog.version_label.text()
    assert dialog.github_link.username == GITHUB_USER
    assert dialog.github_link.url.toString() == f"https://github.com/{GITHUB_USER}"
    labels = [w.text() for w in dialog.findChildren(type(dialog.version_label))]
    assert any(DEVELOPER_NAME in t for t in labels)


def test_invert_and_kind_change_paths(app, qtbot):
    from PyQt6.QtCore import Qt
    from models.relations_table_model import COL_ACTIONS, COL_KIND, INVERT_OPTION, ROLE_IS_MUTUAL

    project, window, controller, table_model = app
    _add(window, "A", UiKind.REFERENCES, "B")
    rid = project.relations()[0].id
    a, b = project.section_by_code("A"), project.section_by_code("B")

    def direction():
        r = project.relation(rid)
        return (project.section(r.source_id).code, project.section(r.target_id).code, r.kind)

    assert direction() == ("A", "B", RelationKind.REF)
    # 1) opción «Invertir» de la celda Relación
    table_model.setData(table_model.index(0, COL_KIND), INVERT_OPTION, Qt.ItemDataRole.EditRole)
    assert direction() == ("B", "A", RelationKind.REF)
    assert "antes" in window.statusBar().currentMessage()
    # 2) icono ⇄ de la columna de acciones
    window.table_view.invertRequested.emit(rid)
    assert direction() == ("A", "B", RelationKind.REF)
    # 3) tecla R con la flecha seleccionada en el mapa
    edge = window.scene.edges[rid]
    window.scene.clearSelection()
    edge.setSelected(True)
    qtbot.keyClick(window.view, Qt.Key.Key_R)
    assert direction() == ("B", "A", RelationKind.REF)
    # 4) convertir en mutua y volver a dirigida con sentido elegido
    controller.relations.set_kind(rid, UiKind.MUTUAL)
    assert direction()[2] is RelationKind.MUTUAL
    assert table_model.index(0, COL_ACTIONS).data(ROLE_IS_MUTUAL) is True
    controller.relations.invert(rid)  # no aplica a mutua: sin cambio
    assert direction()[2] is RelationKind.MUTUAL
    controller.relations.set_direction(rid, a.id, b.id)
    assert direction() == ("A", "B", RelationKind.REF)
    # El combo de edición no ofrece «← Es referenciada por» en una relación existente
    delegate = window.table_view._kind_delegate
    editor = delegate.createEditor(window, None, table_model.index(0, COL_KIND))
    options = [editor.itemText(i) for i in range(editor.count())]
    assert UiKind.REFERENCED_BY.value not in options and INVERT_OPTION in options
    assert "→" in edge.toolTip() and "Clic derecho" in edge.toolTip()


def test_export_report_action_generates_file(app, tmp_path, monkeypatch, qtbot):
    from openpyxl import load_workbook

    import views.components.report_dialog as rd

    project, window, controller, _t = app
    _add(window, "A", UiKind.REFERENCES, "B")
    out = tmp_path / "reporte.xlsx"

    class FakeDialog:
        DialogCode = rd.ReportDialog.DialogCode

        def __init__(self, suggested, parent=None):
            assert suggested.suffix == ".xlsx" and "reporte_secciones" in suggested.name

        def exec(self):
            return self.DialogCode.Accepted

        def result(self):
            return out, True, False

    monkeypatch.setattr(rd, "ReportDialog", FakeDialog)
    window.act_export_report.trigger()
    # El reporte se genera en un hilo de trabajo: esperar a que termine sin bloquear la interfaz.
    qtbot.waitUntil(lambda: "Reporte generado" in window.statusBar().currentMessage(), timeout=30000)
    assert out.exists()
    wb = load_workbook(out)
    assert "Mapa" in wb.sheetnames and wb["Secciones"].max_row == 3
    assert "Reporte generado" in window.statusBar().currentMessage()


def test_toolbar_has_object_name_for_save_state(app):
    _p, window, _c, _t = app
    assert window.toolbar.objectName()
    assert not window.saveState().isEmpty()


def test_completer_filters_by_code_and_title(app):
    from models.section_completer_model import ROLE_CODE, ROLE_KIND

    project, window, controller, _t = app
    project.add_section("31 23 00", "Excavación")
    project.add_section("33 40 00", "Drenaje pluvial")
    controller.rebuild_completer_later.flush()  # la reconstrucción está coalescida (50 ms)
    proxy = window.entry.picker_a._proxy
    proxy.set_query("drenaje")
    assert proxy.rowCount() == 1 and proxy.index(0, 0).data(ROLE_KIND) == "section"
    proxy.set_query("31 exc")
    assert proxy.rowCount() >= 1
    assert proxy.index(0, 0).data(ROLE_CODE) == "31 23 00"  # las del proyecto van primero
    # Código con espacios mal puestos: se compara compactado.
    proxy.set_query("0330 00")
    codes = [proxy.index(r, 0).data(ROLE_CODE) for r in range(proxy.rowCount())]
    assert "03 30 00" in codes
    proxy.set_query("")
    assert proxy.rowCount() > 8000  # catálogo maestro + secciones del proyecto


def test_strict_picker_exact_code_match_creates_canonical_section(app, qtbot):
    from PyQt6.QtCore import Qt

    project, window, controller, table_model = app
    picker_a, picker_b = window.entry.picker_a, window.entry.picker_b
    picker_a.setText("0330 00")
    picker_a.textEdited.emit("0330 00")
    qtbot.keyClick(picker_a, Qt.Key.Key_Return)
    assert picker_a.value() == ("catalog", "033000")
    picker_b.setText("312300")
    picker_b.textEdited.emit("312300")
    qtbot.keyClick(picker_b, Qt.Key.Key_Return)  # Enter en B con elección válida -> agregar
    assert table_model.rowCount() == 1
    codes = sorted(s.code for s in project.sections())
    assert codes == ["03 30 00", "31 23 00"]
    assert project.section_by_code("31 23 00").title == "Excavation and Fill"


def test_strict_picker_rejects_free_text(app, qtbot):
    from PyQt6.QtCore import Qt

    project, window, _c, table_model = app
    picker_a = window.entry.picker_a
    picker_a.setText("texto que no existe")
    picker_a.textEdited.emit("texto que no existe")
    got: list[str] = []
    picker_a.noMatch.connect(got.append)
    qtbot.keyClick(picker_a, Qt.Key.Key_Return)
    assert got == ["texto que no existe"]
    assert picker_a.value() is None and picker_a.property("invalid") is True
    window.entry.add_button.click()
    assert table_model.rowCount() == 0 and project.sections() == []


def test_catalog_panel_double_click_adds_section_and_marks_it(app):
    project, window, controller, _t = app
    key = "312300"
    idx = window.catalog_tree_model.index_for_key(key)
    assert idx.isValid()
    window.catalog_panel.addRequested.emit(key)
    section = project.section_by_code("31 23 00")
    assert section is not None and section.title == "Excavation and Fill"
    controller.catalog.refresh_keys_later.flush()  # marcado en el árbol coalescido (50 ms)
    assert idx.data(ROLE_IN_PROJECT_TREE) is True
    assert section.id in window.scene.nodes


def test_drop_from_catalog_creates_section_at_position(app):
    from PyQt6.QtCore import QPointF

    project, window, controller, _t = app
    window.view.sectionsDropped.emit(["334000", "033000"], QPointF(500.0, 300.0))
    a = project.section_by_code("33 40 00")
    b = project.section_by_code("03 30 00")
    assert a is not None and b is not None
    pa, pb = project.position(a.id), project.position(b.id)
    assert (pa.x, pa.y) == (500.0, 300.0) and pa.pinned
    assert (pb.x, pb.y) == (500.0, 380.0)


def test_custom_section_dialog_suggests_catalog_entry(app):
    from views.components.section_editor_dialog import SectionEditorDialog

    project, window, _c, _t = app
    record = project.master_record("0330 00")
    assert record is not None and record.code == "03 30 00"
    dialog = SectionEditorDialog(project.categories(), window, code="0330 00", is_new=True,
                                 suggestion=(record.code, record.title))
    assert dialog.suggestion_button is not None
    dialog.suggestion_button.click()
    assert dialog.values()[0] == "03 30 00"
    assert dialog.values()[1] == record.title


def test_import_burst_triggers_few_heavy_refreshes(app, qtbot):
    """Importar N secciones no debe reconstruir el autocompletado ni el análisis N veces."""
    from models.project_io import Tables, apply_tables

    project, window, controller, _t = app
    completer_model = controller.completer_model
    qtbot.wait(120)  # vaciar refrescos pendientes del arranque
    rebuilds_before = completer_model.rebuilds
    analysis_before = controller.analysis.refresh_later.fired
    tables = Tables()
    tables.sections = [{"code": f"{d:02d} 10 00", "title": f"Sección {d}", "category": "", "color": "",
                        "status": "", "progress": "", "responsibles": "", "observations": ""} for d in range(1, 61)]
    tables.relations = [{"a": f"{d:02d} 10 00", "kind": "Hace referencia a →", "b": f"{d + 1:02d} 10 00"}
                        for d in range(1, 60)]
    summary = controller.apply_tables_with_progress(tables)
    assert summary is not None and summary.sections_created == 60 and summary.relations_created == 59
    qtbot.wait(200)
    assert completer_model.rebuilds - rebuilds_before <= 3
    assert controller.analysis.refresh_later.fired - analysis_before <= 3
    assert window.analysis.hubs_table.rowCount() == 50  # el panel lista como máximo 50 concentradoras
    assert len(window.scene.nodes) == 60 and len(window.scene.edges) == 59


def test_picker_defers_filtering_while_typing(app, qtbot):
    _p, window, _c, _t = app
    picker = window.entry.picker_a
    total = picker._proxy.rowCount()
    picker.setText("conc")
    picker.textEdited.emit("conc")
    assert picker._proxy.rowCount() == total  # aún no filtró: espera a que el usuario deje de escribir
    qtbot.waitUntil(lambda: picker._proxy.rowCount() < total, timeout=3000)
    assert picker.match_count() > 0
    picker.flush_filter()
