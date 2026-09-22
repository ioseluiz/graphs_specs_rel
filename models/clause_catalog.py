"""Catálogo de cláusulas del pliego (numeración 4.28.N y subcláusulas 4.28.N.M).

Es un catálogo independiente del MasterFormat: se empaqueta como `assets/data/clausulas.sqlite` (generado desde el
Excel del cliente con `scripts/build_clauses_catalog.py`) y el usuario puede reemplazarlo con su propio Excel
(columnas Numeración | Título | Tipo), que se guarda como copia en %APPDATA%\\SpecRel.

Las cláusulas se convierten en nodos de tipo 'clause': solo se conectan y muestran su etiqueta (sin estatus,
avance ni responsables). En Excel la columna Numeración debe ser texto: una celda numérica `4.10` llega como `4.1`.
"""
from __future__ import annotations

import csv
import re
import sqlite3
from dataclasses import dataclass
from datetime import datetime
from functools import cached_property
from pathlib import Path
from typing import Iterable

from config.settings import CLAUSE_CATALOG_PATH, user_clause_catalog_path
from models.relation_normalizer import code_key, search_key, sort_key

CLAUSES_SQL = """
CREATE TABLE IF NOT EXISTS clauses (
    code_key   TEXT PRIMARY KEY,
    code       TEXT NOT NULL,
    title      TEXT NOT NULL,
    level      INTEGER NOT NULL CHECK (level IN (1, 2)),
    parent_key TEXT
);
CREATE INDEX IF NOT EXISTS ix_clauses_parent ON clauses(parent_key);
CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT NOT NULL);
"""

# Número de página pegado al título al extraer del PDF: "50PAGO FINAL", "37Protección De Recursos".
_PAGE_PREFIX = re.compile(r"^\d{1,3}(?=[A-ZÁÉÍÓÚÑ])")
_WS = re.compile(r"\s+")
_CODE_HEADERS = ("numeracion", "numeración", "numero", "número", "codigo", "código", "code", "clausula", "cláusula")
_TITLE_HEADERS = ("titulo", "título", "title", "descripcion", "descripción", "nombre", "texto")
_TYPE_HEADERS = ("tipo", "type", "nivel", "level", "clase")
KIND_LABELS = {1: "Cláusula", 2: "Subcláusula"}
MIN_RECORDS = 5


class ClauseCatalogError(Exception):
    pass


@dataclass(frozen=True)
class ClauseRecord:
    code_key: str
    code: str
    title: str
    level: int                 # 1 = cláusula, 2 = subcláusula
    parent_key: str | None
    quality: str = "ok"        # compatibilidad con el delegate del árbol (MasterFormat usa ok/review)

    @property
    def label(self) -> str:
        return f"{self.code} - {self.title}" if self.title else self.code

    @property
    def kind_label(self) -> str:
        return KIND_LABELS.get(self.level, "Cláusula")

    @cached_property
    def search(self) -> str:
        return f"{self.code_key.lower()} {search_key(self.code)} {search_key(self.title)}"


# ============================================================================ limpieza y construcción
def clean_clause_title(raw: str) -> tuple[str, bool]:
    """Devuelve (título limpio, se_quitó_número_de_página). Colapsa espacios."""
    text = _WS.sub(" ", str(raw or "")).strip()
    cleaned = _PAGE_PREFIX.sub("", text)
    return cleaned.strip(), cleaned != text


def normalize_clause_code(raw) -> str:
    return _WS.sub("", str(raw or "")).strip().rstrip(".")


def _find_column(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    normalized = [search_key(str(h or "")).strip() for h in headers]
    for i, h in enumerate(normalized):
        if h in candidates:
            return i
    for i, h in enumerate(normalized):
        if any(h.startswith(c) for c in candidates if len(c) > 3):
            return i
    return None


def load_clause_rows(path: Path) -> list[dict[str, str]]:
    """Lee un XLSX (primera hoja) o CSV con encabezados Numeración | Título | Tipo (Tipo opcional)."""
    path = Path(path)
    if not path.exists():
        raise ClauseCatalogError(f"No existe el archivo: {path}")
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise ClauseCatalogError("openpyxl no está instalado.") from exc
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ClauseCatalogError(f"No se pudo abrir el archivo Excel: {exc}") from exc
        try:
            rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    elif suffix == ".csv":
        try:
            with open(path, "r", encoding="utf-8-sig", newline="") as fh:
                sample = fh.read(4096)
                fh.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                rows = [list(r) for r in csv.reader(fh, dialect)]
        except (OSError, UnicodeDecodeError) as exc:
            raise ClauseCatalogError(f"No se pudo leer el CSV: {exc}") from exc
    else:
        raise ClauseCatalogError(f"Formato no soportado: {path.suffix}. Use Excel (.xlsx) o CSV.")
    rows = [r for r in rows if any(v is not None and str(v).strip() for v in r)]
    if not rows:
        raise ClauseCatalogError("El archivo está vacío.")
    headers = ["" if h is None else str(h) for h in rows[0]]
    ci = _find_column(headers, _CODE_HEADERS)
    ti = _find_column(headers, _TITLE_HEADERS)
    ki = _find_column(headers, _TYPE_HEADERS)
    if ci is None:
        raise ClauseCatalogError("No se encontró la columna «Numeración» (también se acepta Número o Código).")
    if ti is None:
        ti = 1 if ci != 1 and len(headers) > 1 else None
    out: list[dict[str, str]] = []
    for r in rows[1:]:
        cell = r[ci] if ci < len(r) else None
        if cell is None or not str(cell).strip():
            continue
        out.append({
            "code": str(cell),
            "title": "" if ti is None or ti >= len(r) or r[ti] is None else str(r[ti]),
            "type": "" if ki is None or ki >= len(r) or r[ki] is None else str(r[ki]),
        })
    return out


def build_clause_records(rows: Iterable[dict[str, str]]) -> tuple[list[ClauseRecord], dict[str, int]]:
    """Limpia, deduplica y jerarquiza. Devuelve (registros, estadísticas)."""
    stats = {"rows": 0, "cleaned_titles": 0, "duplicates": 0, "skipped": 0}
    by_key: dict[str, ClauseRecord] = {}
    for row in rows:
        stats["rows"] += 1
        code = normalize_clause_code(row.get("code", ""))
        if not code or not re.fullmatch(r"\d+(?:\.\d+)+", code):
            stats["skipped"] += 1
            continue
        title, cleaned = clean_clause_title(row.get("title", ""))
        stats["cleaned_titles"] += int(cleaned)
        kind_text = search_key(row.get("type", "")).strip()
        if kind_text.startswith("sub"):
            level = 2
        elif kind_text.startswith("clau"):
            level = 1
        else:
            level = 2 if code.count(".") >= 3 else 1
        parent_key = code_key(code.rsplit(".", 1)[0]) if level == 2 else None
        key = code_key(code)
        if key in by_key:
            stats["duplicates"] += 1
        by_key[key] = ClauseRecord(key, code, title, level, parent_key)
    records = sorted(by_key.values(), key=lambda r: sort_key(r.code_key))
    return records, stats


def clauses_root_label(records: Iterable[ClauseRecord]) -> str:
    """«Cláusulas 4.28» si todas las cláusulas comparten prefijo; si no, «Cláusulas»."""
    prefixes = {r.code.rsplit(".", 1)[0] for r in records if r.level == 1 and "." in r.code}
    return f"Cláusulas {prefixes.pop()}" if len(prefixes) == 1 else "Cláusulas"


def write_clauses_sqlite(records: Iterable[ClauseRecord], path: Path, meta: dict[str, str] | None = None) -> int:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    if tmp.exists():
        tmp.unlink()
    conn = sqlite3.connect(tmp)
    try:
        conn.executescript(CLAUSES_SQL)
        rows = [(r.code_key, r.code, r.title, r.level, r.parent_key) for r in records]
        conn.executemany("INSERT INTO clauses VALUES (?, ?, ?, ?, ?)", rows)
        meta = dict(meta or {})
        meta.setdefault("built_at", datetime.now().isoformat(timespec="seconds"))
        meta["rows"] = str(len(rows))
        conn.executemany("INSERT OR REPLACE INTO meta (key, value) VALUES (?, ?)", list(meta.items()))
        conn.commit()
    finally:
        conn.close()
    if path.exists():
        path.unlink()
    tmp.replace(path)
    return len(rows)


# ============================================================================ consulta
class ClauseCatalog:
    def __init__(self, path: Path | None = None) -> None:
        self.path: Path | None = None
        self.meta: dict[str, str] = {}
        self._records: dict[str, ClauseRecord] = {}
        self._children: dict[str | None, list[ClauseRecord]] = {}
        self.load(path)

    # ------------------------------------------------------------------ carga
    @staticmethod
    def resolve_path() -> Path | None:
        user = user_clause_catalog_path()
        if user.exists():
            return user
        if CLAUSE_CATALOG_PATH.exists():
            return CLAUSE_CATALOG_PATH
        return None

    def load(self, path: Path | None = None) -> None:
        self.path = path if path is not None else self.resolve_path()
        self._records, self._children, self.meta = {}, {}, {}
        if self.path is None or not self.path.exists():
            return
        try:
            conn = sqlite3.connect(f"file:{Path(self.path).as_posix()}?mode=ro", uri=True)
        except sqlite3.Error as exc:
            raise ClauseCatalogError(f"No se pudo abrir el catálogo de cláusulas: {exc}") from exc
        try:
            for row in conn.execute("SELECT code_key, code, title, level, parent_key FROM clauses"):
                rec = ClauseRecord(*row)
                self._records[rec.code_key] = rec
            self.meta = {k: v for k, v in conn.execute("SELECT key, value FROM meta")}
        except sqlite3.Error as exc:
            raise ClauseCatalogError(f"El catálogo de cláusulas no tiene el formato esperado: {exc}") from exc
        finally:
            conn.close()
        for rec in sorted(self._records.values(), key=lambda r: sort_key(r.code_key)):
            parent = rec.parent_key if rec.parent_key in self._records else None
            self._children.setdefault(parent, []).append(rec)

    @property
    def available(self) -> bool:
        return bool(self._records)

    @property
    def is_user_copy(self) -> bool:
        try:
            return self.path is not None and Path(self.path).resolve() == user_clause_catalog_path().resolve()
        except OSError:
            return False

    def __len__(self) -> int:
        return len(self._records)

    @property
    def root_label(self) -> str:
        return clauses_root_label(self._records.values())

    # ------------------------------------------------------------------ consultas
    def all(self) -> list[ClauseRecord]:
        return sorted(self._records.values(), key=lambda r: sort_key(r.code_key))

    def get(self, code_or_key: str) -> ClauseRecord | None:
        key = code_key(normalize_clause_code(code_or_key))
        return self._records.get(key)

    def roots(self) -> list[ClauseRecord]:
        return list(self._children.get(None, []))

    def children(self, parent_key: str) -> list[ClauseRecord]:
        return list(self._children.get(parent_key, []))

    def search(self, text: str, limit: int = 50) -> list[ClauseRecord]:
        query = search_key(text).strip()
        if not query:
            return self.all()[:limit]
        tokens = query.split()
        out = [r for r in self.all() if all(tok in r.search for tok in tokens)]
        return out[:limit]

    # ------------------------------------------------------------------ reemplazo por el usuario
    @staticmethod
    def build_user_catalog_file(source: Path, target: Path | None = None) -> tuple[Path, int]:
        """Lee, limpia y escribe la copia del usuario. Función pura: apta para un hilo.

        Devuelve (ruta escrita, número de cláusulas). No toca la instancia: llamar después a `load`.
        """
        records, _stats = build_clause_records(load_clause_rows(Path(source)))
        if len(records) < MIN_RECORDS:
            raise ClauseCatalogError(
                "El archivo tiene muy pocas cláusulas válidas. Compruebe que la columna «Numeración» sea texto "
                "(por ejemplo 4.28.3.1) y que exista una columna «Título».")
        target = Path(target) if target is not None else user_clause_catalog_path()
        count = write_clauses_sqlite(records, target, {"source": Path(source).name, "edition": "usuario"})
        return target, count

    def replace_from_file(self, source: Path) -> int:
        target, count = self.build_user_catalog_file(source)
        self.load(target)
        return count

    def reset_to_bundled(self) -> None:
        user = user_clause_catalog_path()
        if user.exists():
            user.unlink()
        self.load(CLAUSE_CATALOG_PATH if CLAUSE_CATALOG_PATH.exists() else None)
