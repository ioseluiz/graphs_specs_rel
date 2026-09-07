"""Catálogo maestro MasterFormat: construcción (limpieza y clasificación), SQLite y consulta en memoria.

El catálogo empaquetado se genera con `scripts/build_master_catalog.py` a partir del Excel del
cliente. El usuario puede reemplazarlo con su propio archivo (Excel/CSV); esa copia se guarda en
%APPDATA%\\SpecRel y tiene prioridad sobre la empaquetada. Las correcciones de clasificación del
usuario se guardan aparte (JSON) y se aplican sobre cualquier catálogo.
"""
from __future__ import annotations

import csv
import json
import re
import sqlite3
from dataclasses import dataclass, replace
from functools import cached_property
from datetime import datetime
from pathlib import Path
from typing import Iterable, Iterator

from config.settings import (
    BASE_DIR,
    MASTER_CATALOG_PATH,
    catalog_edits_path,
    category_overrides_path,
    user_catalog_path,
)
from models.relation_normalizer import code_key, search_key

REVIEW_TITLE = "(título por verificar)"
MAX_TITLE_LEN = 80
DEFAULT_RULES_PATH = BASE_DIR / "scripts" / "category_rules.csv"
_INNER_CODE = re.compile(r"\b\d{2} \d{2} \d{2}\b")
_KEYWORD_INDEX = re.compile(r"\bkeyword index\b", re.IGNORECASE)
_WS = re.compile(r"\s+")

CATALOG_SQL = """
CREATE TABLE IF NOT EXISTS catalog (
    code_key   TEXT PRIMARY KEY,
    code       TEXT NOT NULL,
    title_en   TEXT NOT NULL,
    title_es   TEXT,
    level      INTEGER NOT NULL,
    parent_key TEXT,
    division   TEXT NOT NULL,
    quality    TEXT NOT NULL CHECK (quality IN ('ok', 'review')),
    category   TEXT
);
CREATE INDEX IF NOT EXISTS ix_catalog_parent ON catalog(parent_key);
CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT);
"""


class MasterCatalogError(Exception):
    pass


@dataclass(frozen=True)
class CatalogRecord:
    code_key: str
    code: str
    title_en: str
    title_es: str | None
    level: int
    parent_key: str | None
    division: str
    quality: str  # 'ok' | 'review'
    category: str | None = None  # clasificación por defecto (nombre de categoría)

    @property
    def title(self) -> str:
        return (self.title_es or "").strip() or self.title_en

    @property
    def label(self) -> str:
        return f"{self.code} - {self.title}" if self.title else self.code

    @cached_property
    def search(self) -> str:
        """Texto de búsqueda normalizado; se calcula una sola vez por registro (antes, en cada tecla)."""
        return f"{self.code_key.lower()} {search_key(self.code)} {search_key(self.title_en)} {search_key(self.title_es or '')}"


# ============================================================================ limpieza
def normalize_code(raw: str) -> str:
    """'0330 00' -> '03 30 00'; '02 03 01.19' se conserva; otros formatos se compactan en espacios."""
    text = _WS.sub(" ", str(raw or "")).strip()
    digits_only = re.sub(r"[^0-9.]", "", text)
    if re.fullmatch(r"\d{6}(\.\d{2})?", digits_only):
        base = f"{digits_only[0:2]} {digits_only[2:4]} {digits_only[4:6]}"
        return base + digits_only[6:] if len(digits_only) > 6 else base
    if re.fullmatch(r"\d{2} ?\d{2} ?\d{2}(\.\d{2})?", text):
        compact = text.replace(" ", "")
        return f"{compact[0:2]} {compact[2:4]} {compact[4:6]}" + compact[6:]
    return text


def infer_level(code: str) -> int:
    key = code_key(code)
    if "." in code or len(key) > 6:
        return 4
    if key.endswith("0000"):
        return 1
    if key.endswith("00"):
        return 2
    return 3


def derive_parent_key(code: str, level: int) -> str | None:
    key = code_key(code)
    if level <= 1 or len(key) < 6:
        return None
    if level == 4:
        return key[:6]
    if level == 3:
        return key[:4] + "00"
    return key[:2] + "0000"


def clean_title(raw: str) -> tuple[str, str]:
    """Devuelve (título, calidad). Corta ruido de extracción del PDF."""
    text = str(raw or "")
    text = text.split("\n", 1)[0]
    match = _INNER_CODE.search(text)
    if match and match.start() > 0:
        text = text[: match.start()]
    kw = _KEYWORD_INDEX.search(text)
    if kw:
        text = text[: kw.start()]
    text = _WS.sub(" ", text).strip(" .,;:-–—")
    # Los títulos MasterFormat empiezan en mayúscula; un inicio en minúscula delata texto de índice.
    starts_lower = bool(text) and text[0].isalpha() and text[0].islower()
    if not text or len(text) > MAX_TITLE_LEN or (match and match.start() == 0) or starts_lower:
        return REVIEW_TITLE, "review"
    return text, "ok"


def build_records(rows: Iterable[dict]) -> list[CatalogRecord]:
    """rows: dicts con claves code, title_en, title_es (opc.), level (opc.), category (opc.)."""
    records: dict[str, CatalogRecord] = {}
    for row in rows:
        code = normalize_code(row.get("code", ""))
        key = code_key(code)
        if not key or not re.search(r"\d", key):
            continue
        title_en, quality = clean_title(row.get("title_en", ""))
        title_es = _WS.sub(" ", str(row.get("title_es") or "")).strip() or None
        category = _WS.sub(" ", str(row.get("category") or "")).strip() or None
        try:
            level = int(row.get("level") or 0)
        except (TypeError, ValueError):
            level = 0
        if level not in (1, 2, 3, 4):
            level = infer_level(code)
        record = CatalogRecord(
            code_key=key, code=code, title_en=title_en, title_es=title_es, level=level,
            parent_key=derive_parent_key(code, level), division=key[:2], quality=quality, category=category,
        )
        existing = records.get(key)
        if existing is None or (existing.quality == "review" and quality == "ok") or (
            existing.title_es is None and title_es
        ):
            records[key] = record
    return sorted(records.values(), key=lambda r: r.code_key)


def load_title_overrides(path: Path) -> dict[str, str]:
    """CSV 'code;title_en' con títulos corregidos a mano (para códigos cuyo título llegó dañado)."""
    if not path.exists():
        return {}
    overrides: dict[str, str] = {}
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh, delimiter=";")
        next(reader, None)
        for row in reader:
            if len(row) >= 2 and row[0].strip() and row[1].strip():
                overrides[code_key(normalize_code(row[0]))] = _WS.sub(" ", row[1]).strip()
    return overrides


def apply_title_overrides(records: list[CatalogRecord], overrides: dict[str, str]) -> list[CatalogRecord]:
    out: list[CatalogRecord] = []
    for rec in records:
        title = overrides.get(rec.code_key)
        if title:
            rec = replace(rec, title_en=title, quality="ok")
        out.append(rec)
    return out


# ============================================================================ clasificación
def load_category_rules(path: Path = DEFAULT_RULES_PATH) -> list[tuple[str, str]]:
    """CSV 'prefix;category'. Prefijo vacío = regla por defecto; '??' = cualquier división."""
    if not path.exists():
        return []
    rules: list[tuple[str, str]] = []
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh, delimiter=";")
        next(reader, None)
        for row in reader:
            if not row or row[0].strip().startswith("#") or len(row) < 2:
                continue
            prefix = re.sub(r"[^0-9?]", "", row[0]).upper()
            category = _WS.sub(" ", row[1]).strip()
            if category:
                rules.append((prefix, category))
    return rules


def match_category(key: str, rules: Iterable[tuple[str, str]]) -> str | None:
    """Regla con el prefijo más largo que coincide con la clave ('??' comodín de división)."""
    best: tuple[int, str] | None = None
    for prefix, category in rules:
        if prefix.startswith("??"):
            candidate = key[:2] + prefix[2:]
        else:
            candidate = prefix
        if key.startswith(candidate) and (best is None or len(candidate) > best[0]):
            best = (len(candidate), category)
    return best[1] if best else None


def classify(records: list[CatalogRecord], rules: list[tuple[str, str]], overwrite: bool = False) -> list[CatalogRecord]:
    """Asigna categoría por reglas a los registros sin categoría (o a todos si overwrite)."""
    if not rules:
        return records
    out: list[CatalogRecord] = []
    for rec in records:
        if rec.category and not overwrite:
            out.append(rec)
            continue
        out.append(replace(rec, category=match_category(rec.code_key, rules)))
    return out


# ============================================================================ lectura de archivos fuente
_CODE_HEADERS = ("codigo_display", "codigo", "código", "code", "numero", "número", "seccion", "sección")
_TITLE_EN_HEADERS = ("nombre", "title_en", "title", "titulo", "título", "descripcion", "descripción", "name")
_TITLE_ES_HEADERS = ("title_es", "titulo_es", "título es", "titulo es", "descripcion_es", "nombre_es", "español", "espanol")
_LEVEL_HEADERS = ("nivel", "level")
_CATEGORY_HEADERS = ("categoria", "categoría", "category", "clasificacion", "clasificación")


def _pick(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    norm = [search_key(h) for h in headers]
    for cand in candidates:
        c = search_key(cand)
        if c in norm:
            return norm.index(c)
    return None


def load_source_rows(path: Path) -> list[dict]:
    """Lee Excel (hoja con encabezados reconocibles) o CSV y devuelve filas normalizadas."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise MasterCatalogError("openpyxl no está instalado.") from exc
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise MasterCatalogError(f"No se pudo abrir el Excel: {exc}") from exc
        try:
            for ws in wb.worksheets:
                it = ws.iter_rows(values_only=True)
                first = next(it, None)
                if not first:
                    continue
                headers = ["" if h is None else str(h) for h in first]
                if _pick(headers, _CODE_HEADERS) is None:
                    continue
                return list(_map_rows(headers, it))
        finally:
            wb.close()
        raise MasterCatalogError("Ninguna hoja tiene una columna de código reconocible "
                                 "(codigo_display / Número / code).")
    if suffix == ".csv":
        try:
            with open(path, "r", encoding="utf-8-sig", newline="") as fh:
                sample = fh.read(4096)
                fh.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.reader(fh, dialect)
                headers = [h.strip() for h in next(reader, [])]
                if _pick(headers, _CODE_HEADERS) is None:
                    raise MasterCatalogError("El CSV no tiene una columna de código reconocible.")
                return list(_map_rows(headers, reader))
        except (OSError, UnicodeDecodeError) as exc:
            raise MasterCatalogError(f"No se pudo leer el CSV: {exc}") from exc
    if suffix in (".sqlite", ".db", ".sqlite3"):
        return load_catalog_sqlite_rows(path)
    raise MasterCatalogError(f"Formato no soportado: {suffix}")


def _map_rows(headers: list[str], rows: Iterable) -> Iterator[dict]:
    ci = _pick(headers, _CODE_HEADERS)
    ti = _pick(headers, _TITLE_EN_HEADERS)
    ei = _pick(headers, _TITLE_ES_HEADERS)
    li = _pick(headers, _LEVEL_HEADERS)
    gi = _pick(headers, _CATEGORY_HEADERS)
    for row in rows:
        values = list(row)
        get = lambda i: (values[i] if i is not None and i < len(values) else None)  # noqa: E731
        if not get(ci):
            continue
        yield {"code": get(ci), "title_en": get(ti) or "", "title_es": get(ei), "level": get(li),
               "category": get(gi)}


def _table_columns(conn: sqlite3.Connection, table: str) -> set[str]:
    return {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}


def load_catalog_sqlite_rows(path: Path) -> list[dict]:
    """Filas desde un SQLite con tabla `catalog` (este formato) o `master_format` (02_master_format_database)."""
    try:
        conn = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True)
    except sqlite3.Error as exc:
        raise MasterCatalogError(f"No se pudo abrir la base: {exc}") from exc
    try:
        tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if "catalog" in tables:
            has_cat = "category" in _table_columns(conn, "catalog")
            cols = "code, title_en, title_es, level" + (", category" if has_cat else ", NULL")
            cur = conn.execute(f"SELECT {cols} FROM catalog")
            return [{"code": c, "title_en": t, "title_es": e, "level": lv, "category": g} for c, t, e, lv, g in cur]
        if "master_format" in tables:
            cur = conn.execute("SELECT codigo_display, nombre, nivel FROM master_format")
            return [{"code": c, "title_en": t, "title_es": None, "level": lv, "category": None} for c, t, lv in cur]
        raise MasterCatalogError("La base no contiene una tabla 'catalog' ni 'master_format'.")
    except sqlite3.Error as exc:
        raise MasterCatalogError(str(exc)) from exc
    finally:
        conn.close()


# ============================================================================ escritura
def write_catalog_sqlite(records: Iterable[CatalogRecord], path: Path, meta: dict[str, str] | None = None) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    if tmp.exists():
        tmp.unlink()
    conn = sqlite3.connect(tmp)
    try:
        conn.executescript(CATALOG_SQL)
        rows = [(r.code_key, r.code, r.title_en, r.title_es, r.level, r.parent_key, r.division, r.quality, r.category)
                for r in records]
        conn.executemany("INSERT INTO catalog VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)", rows)
        meta = dict(meta or {})
        meta.setdefault("built_at", datetime.now().isoformat(timespec="seconds"))
        meta["rows"] = str(len(rows))
        conn.executemany("INSERT OR REPLACE INTO meta (key, value) VALUES (?, ?)", list(meta.items()))
        conn.commit()
    finally:
        conn.close()
    if path.exists():
        path.unlink()
    tmp.replace(path)
    return len(rows)


# ============================================================================ consulta
class MasterCatalog:
    """Catálogo cargado en memoria. `path=None` elige la copia del usuario o la empaquetada."""

    def __init__(self, path: Path | None = None, overrides_path: Path | None = None,
                 edits_path: Path | None = None) -> None:
        self.path: Path | None = None
        self.meta: dict[str, str] = {}
        self._base: dict[str, CatalogRecord] = {}       # tal como viene del archivo
        self._records: dict[str, CatalogRecord] = {}    # visibles, con ediciones del usuario aplicadas
        self._hidden: dict[str, CatalogRecord] = {}     # ocultas por el usuario
        self._children: dict[str | None, list[CatalogRecord]] = {}
        self._overrides_path = overrides_path if overrides_path is not None else category_overrides_path()
        self._edits_path = edits_path if edits_path is not None else catalog_edits_path()
        self._code_overrides: dict[str, str] = {}
        self._prefix_overrides: dict[str, str] = {}
        self._edits: dict = {"titles": {}, "added": {}, "hidden": []}
        self._load_overrides()
        self._load_edits()
        self.load(path)

    # ------------------------------------------------------------------ carga
    @staticmethod
    def resolve_path() -> Path | None:
        user = user_catalog_path()
        if user.exists():
            return user
        if MASTER_CATALOG_PATH.exists():
            return MASTER_CATALOG_PATH
        return None

    def load(self, path: Path | None = None) -> None:
        self.path = path if path is not None else self.resolve_path()
        self._base, self._records, self._hidden, self._children, self.meta = {}, {}, {}, {}, {}
        if self.path is None or not self.path.exists():
            self._apply_edits()
            return
        try:
            conn = sqlite3.connect(f"file:{self.path.as_posix()}?mode=ro", uri=True)
        except sqlite3.Error as exc:
            raise MasterCatalogError(f"No se pudo abrir el catálogo: {exc}") from exc
        try:
            has_cat = "category" in _table_columns(conn, "catalog")
            cols = "code_key, code, title_en, title_es, level, parent_key, division, quality" + \
                   (", category" if has_cat else ", NULL")
            for row in conn.execute(f"SELECT {cols} FROM catalog ORDER BY code_key"):
                rec = CatalogRecord(*row)
                self._base[rec.code_key] = rec
            self.meta = {k: v for k, v in conn.execute("SELECT key, value FROM meta")}
        except sqlite3.Error as exc:
            raise MasterCatalogError(f"El catálogo no tiene el formato esperado: {exc}") from exc
        finally:
            conn.close()
        if not has_cat or not any(r.category for r in self._base.values()):
            # Copias antiguas sin clasificación: aplicar las reglas incluidas en la aplicación.
            rules = load_category_rules()
            if rules:
                for key, rec in list(self._base.items()):
                    self._base[key] = replace(rec, category=match_category(key, rules))
        self._apply_edits()

    def _apply_edits(self) -> None:
        """Construye las vistas visibles/ocultas a partir del archivo + ediciones del usuario."""
        records: dict[str, CatalogRecord] = dict(self._base)
        for key, data in self._edits.get("added", {}).items():
            if key in records:
                continue
            try:
                records[key] = CatalogRecord(
                    code_key=key, code=data["code"], title_en=data.get("title_en") or "",
                    title_es=data.get("title_es") or None, level=int(data.get("level") or infer_level(data["code"])),
                    parent_key=data.get("parent_key"), division=key[:2], quality="ok",
                    category=data.get("category"),
                )
            except (KeyError, ValueError):
                continue
        for key, data in self._edits.get("titles", {}).items():
            rec = records.get(key)
            if rec is not None:
                records[key] = replace(rec, title_en=data.get("title_en") or rec.title_en,
                                       title_es=data.get("title_es") or None, quality="ok")
        hidden = set(self._edits.get("hidden", []))
        self._hidden = {k: r for k, r in records.items() if k in hidden}
        self._records = {k: r for k, r in sorted(records.items()) if k not in hidden}
        self._children = {}
        for rec in self._records.values():
            self._children.setdefault(rec.parent_key, []).append(rec)

    @property
    def is_user_copy(self) -> bool:
        return self.path is not None and self.path == user_catalog_path()

    @property
    def available(self) -> bool:
        return bool(self._records)

    def __len__(self) -> int:
        return len(self._records)

    # ------------------------------------------------------------------ consultas
    def all(self, include_hidden: bool = False) -> list[CatalogRecord]:
        if include_hidden:
            merged = {**self._records, **self._hidden}
            return [merged[k] for k in sorted(merged)]
        return list(self._records.values())

    def get(self, code_or_key: str, include_hidden: bool = False) -> CatalogRecord | None:
        key = code_key(code_or_key)
        rec = self._records.get(key)
        if rec is None and include_hidden:
            rec = self._hidden.get(key)
        return rec

    # ------------------------------------------------------------------ ediciones del usuario
    def is_hidden(self, code_key_: str) -> bool:
        return code_key_ in self._hidden

    def is_user_added(self, code_key_: str) -> bool:
        return code_key_ in self._edits.get("added", {})

    def is_title_edited(self, code_key_: str) -> bool:
        return code_key_ in self._edits.get("titles", {})

    @property
    def edit_count(self) -> int:
        e = self._edits
        return len(e.get("titles", {})) + len(e.get("added", {})) + len(e.get("hidden", []))

    def set_title(self, code_key_: str, title_en: str, title_es: str | None = None) -> None:
        title_en = _WS.sub(" ", title_en or "").strip()
        if not title_en:
            raise MasterCatalogError("El título no puede estar vacío.")
        key = code_key(code_key_)
        base = self._base.get(key)
        added = self._edits["added"].get(key)
        if added is not None:
            added["title_en"], added["title_es"] = title_en, (title_es or None)
        elif base is not None and base.title_en == title_en and (base.title_es or None) == (title_es or None):
            self._edits["titles"].pop(key, None)  # volvió al original
        else:
            self._edits["titles"][key] = {"title_en": title_en, "title_es": (title_es or None)}
        self._save_edits()
        self._apply_edits()

    def add_section(self, code: str, title_en: str, title_es: str | None = None,
                    category: str | None = None) -> CatalogRecord:
        """Agrega una sección al catálogo (edición del usuario). El padre se deriva del código."""
        code = normalize_code(code)
        key = code_key(code)
        if not key or not re.search(r"\d", key):
            raise MasterCatalogError("El número de sección no es válido.")
        if key in self._records or key in self._hidden:
            raise MasterCatalogError(f"La sección {code} ya existe en el catálogo.")
        title_en = _WS.sub(" ", title_en or "").strip()
        if not title_en:
            raise MasterCatalogError("Indique la descripción de la sección.")
        level = infer_level(code)
        parent_key = derive_parent_key(code, level)
        if category is None:
            category = match_category(key, load_category_rules())
        self._edits["added"][key] = {"code": code, "title_en": title_en, "title_es": title_es or None,
                                     "level": level, "parent_key": parent_key, "category": category}
        self._save_edits()
        self._apply_edits()
        return self._records[key]

    def hide(self, code_key_: str) -> None:
        key = code_key(code_key_)
        if key in self._edits["added"]:
            del self._edits["added"][key]  # una sección agregada por el usuario se elimina del todo
            self._edits["titles"].pop(key, None)
        elif key not in self._edits["hidden"]:
            self._edits["hidden"].append(key)
        self._save_edits()
        self._apply_edits()

    def unhide(self, code_key_: str) -> None:
        key = code_key(code_key_)
        if key in self._edits["hidden"]:
            self._edits["hidden"].remove(key)
            self._save_edits()
            self._apply_edits()

    def restore_entry(self, code_key_: str) -> None:
        """Quita cualquier edición del usuario sobre esa sección (título, ocultamiento, agregado)."""
        key = code_key(code_key_)
        self._edits["titles"].pop(key, None)
        self._edits["added"].pop(key, None)
        if key in self._edits["hidden"]:
            self._edits["hidden"].remove(key)
        self._save_edits()
        self._apply_edits()

    def clear_edits(self) -> None:
        self._edits = {"titles": {}, "added": {}, "hidden": []}
        self._save_edits()
        self._apply_edits()

    def _load_edits(self) -> None:
        self._edits = {"titles": {}, "added": {}, "hidden": []}
        path = self._edits_path
        if path is None or not path.exists():
            return
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            self._edits["titles"] = dict(data.get("titles") or {})
            self._edits["added"] = dict(data.get("added") or {})
            self._edits["hidden"] = [str(k) for k in (data.get("hidden") or [])]
        except (OSError, ValueError):
            pass

    def _save_edits(self) -> None:
        path = self._edits_path
        if path is None:
            return
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps(self._edits, ensure_ascii=False, indent=2), encoding="utf-8")
        except OSError:
            pass

    # ------------------------------------------------------------------ exportación
    def export_xlsx(self, path: Path, include_hidden: bool = False) -> int:
        """Excel editable con el mismo formato que acepta «Reemplazar catálogo…»."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:  # pragma: no cover
            raise MasterCatalogError("openpyxl no está instalado.") from exc
        wb = Workbook()
        ws = wb.active
        ws.title = "master_format"
        headers = ["Número", "Descripción", "title_es", "Nivel", "Categoría", "Calidad", "Oculta"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
        for rec in self.all(include_hidden=include_hidden):
            ws.append([rec.code, rec.title_en, rec.title_es or "", rec.level,
                       self.effective_category(rec) or "", rec.quality, "sí" if self.is_hidden(rec.code_key) else ""])
        for col, width in zip("ABCDEFG", (14, 60, 40, 8, 26, 10, 8)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        try:
            wb.save(path)
        except OSError as exc:
            raise MasterCatalogError(f"No se pudo guardar el archivo: {exc}") from exc
        return ws.max_row - 1

    def children(self, parent_key: str | None) -> list[CatalogRecord]:
        return list(self._children.get(parent_key, []))

    def divisions(self) -> list[CatalogRecord]:
        return [r for r in self._children.get(None, []) if r.level == 1]

    def search(self, text: str, limit: int = 50) -> list[CatalogRecord]:
        tokens = search_key(text).split()
        if not tokens:
            return []
        out: list[CatalogRecord] = []
        for rec in self._records.values():
            hay = rec.search
            if all(t in hay for t in tokens):
                out.append(rec)
                if len(out) >= limit:
                    break
        return out

    def suggest_for_code(self, raw: str) -> CatalogRecord | None:
        """Sugerencia por código tipeado con errores de espacios: '0330 00' -> 03 30 00."""
        return self._records.get(code_key(normalize_code(raw)))

    # ------------------------------------------------------------------ clasificación efectiva
    def effective_category(self, record: CatalogRecord | str) -> str | None:
        """Código corregido por el usuario › prefijo corregido más largo › categoría del catálogo."""
        rec = record if isinstance(record, CatalogRecord) else self.get(record)
        if rec is None:
            return None
        if rec.code_key in self._code_overrides:
            return self._code_overrides[rec.code_key]
        best: tuple[int, str] | None = None
        for prefix, category in self._prefix_overrides.items():
            if rec.code_key.startswith(prefix) and (best is None or len(prefix) > best[0]):
                best = (len(prefix), category)
        if best is not None:
            return best[1]
        return rec.category

    def has_override(self, code_key_: str) -> bool:
        """True si el usuario corrigió esta sección (por código) o su rama (prefijo propio)."""
        if code_key_ in self._code_overrides:
            return True
        rec = self._records.get(code_key_)
        level = rec.level if rec else infer_level(code_key_)
        return _branch_prefix(code_key_, level) in self._prefix_overrides

    def set_category_override(self, code_key_: str, category: str, whole_branch: bool = False) -> None:
        """Corrección del usuario: una sección, o toda la rama (prefijo) que cuelga de ella."""
        rec = self._records.get(code_key_)
        if whole_branch:
            prefix = _branch_prefix(code_key_, rec.level if rec else infer_level(code_key_))
            self._prefix_overrides[prefix] = category
            # Las correcciones de códigos dentro de la rama quedan obsoletas.
            for key in [k for k in self._code_overrides if k.startswith(prefix)]:
                del self._code_overrides[key]
        else:
            self._code_overrides[code_key_] = category
        self._save_overrides()

    def clear_category_override(self, code_key_: str) -> None:
        rec = self._records.get(code_key_)
        self._code_overrides.pop(code_key_, None)
        if rec is not None:
            self._prefix_overrides.pop(_branch_prefix(code_key_, rec.level), None)
        self._save_overrides()

    def clear_category_overrides(self) -> None:
        self._code_overrides.clear()
        self._prefix_overrides.clear()
        self._save_overrides()

    @property
    def override_count(self) -> int:
        return len(self._code_overrides) + len(self._prefix_overrides)

    def _load_overrides(self) -> None:
        self._code_overrides, self._prefix_overrides = {}, {}
        path = self._overrides_path
        if path is None or not path.exists():
            return
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            self._code_overrides = {str(k): str(v) for k, v in (data.get("codes") or {}).items()}
            self._prefix_overrides = {str(k): str(v) for k, v in (data.get("prefixes") or {}).items()}
        except (OSError, ValueError):
            pass

    def _save_overrides(self) -> None:
        path = self._overrides_path
        if path is None:
            return
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps({"codes": self._code_overrides, "prefixes": self._prefix_overrides},
                                       ensure_ascii=False, indent=2), encoding="utf-8")
        except OSError:
            pass

    # ------------------------------------------------------------------ reemplazo por el usuario
    @staticmethod
    def build_user_catalog_file(source: Path, target: Path | None = None) -> tuple[Path, int]:
        """Lee, limpia, clasifica y escribe la copia del usuario. Función pura: apta para un hilo.

        Devuelve (ruta escrita, número de secciones). No toca la instancia: llamar después a `load`.
        """
        rows = load_source_rows(source)
        records = classify(build_records(rows), load_category_rules())
        if len(records) < 10:
            raise MasterCatalogError("El archivo tiene muy pocas secciones válidas para ser un catálogo.")
        target = target if target is not None else user_catalog_path()
        count = write_catalog_sqlite(records, target, {"source": source.name, "edition": "usuario"})
        return target, count

    def replace_from_file(self, source: Path) -> int:
        target, count = self.build_user_catalog_file(source)
        self.load(target)
        return count

    def reset_to_bundled(self) -> None:
        user = user_catalog_path()
        if user.exists():
            user.unlink()
        self.load(MASTER_CATALOG_PATH if MASTER_CATALOG_PATH.exists() else None)


def _branch_prefix(key: str, level: int) -> str:
    """Prefijo que abarca la rama de una sección: división (2), nivel 2 (4), nivel 3 (6), nivel 4 (clave)."""
    if level <= 1:
        return key[:2]
    if level == 2:
        return key[:4]
    if level == 3:
        return key[:6]
    return key
