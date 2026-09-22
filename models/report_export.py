"""Reporte de secciones en Excel: resumen, secciones, por responsable, relaciones y mapa opcional.

Función pura sobre ProjectModel (sin Qt) para poder probarla con un proyecto en memoria.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

from config import palette
from models.relation_normalizer import denormalize

if TYPE_CHECKING:
    from models.project_model import ProjectModel

SHEET_SUMMARY, SHEET_SECTIONS, SHEET_BY_RESP, SHEET_RELATIONS, SHEET_MAP = (
    "Resumen", "Secciones", "Por responsable", "Relaciones", "Mapa")
NO_RESPONSIBLE = "(sin responsable)"
HEADER_FILL = "1F4E79"
BAR_COLOR = "4472C4"


class ReportExportError(Exception):
    pass


@dataclass
class ReportStats:
    path: Path
    sections: int = 0
    relations: int = 0
    avg_progress: float = 0.0
    completed: int = 0
    without_responsible: int = 0
    sheets: list[str] = field(default_factory=list)


def _hex(color: str | None, default: str) -> str:
    c = (color or default).lstrip("#")
    return c.upper() if len(c) == 6 else default.lstrip("#").upper()


def export_report_xlsx(model: "ProjectModel", path: Path, map_image: Path | None = None) -> ReportStats:
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import DataBarRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise ReportExportError("openpyxl no está instalado.") from exc

    meta = model.meta()
    sections = model.sections()
    # Las cláusulas del pliego no tienen estatus, avance ni responsables: quedan fuera de los indicadores.
    plain = [s for s in sections if not s.is_clause]
    clauses = [s for s in sections if s.is_clause]
    relations = model.relations()
    now = datetime.now()
    stats = ReportStats(path=path, sections=len(sections), relations=len(relations))
    stats.avg_progress = (sum(s.progress for s in plain) / len(plain)) if plain else 0.0
    stats.completed = sum(1 for s in plain if s.progress >= 100)
    stats.without_responsible = sum(1 for s in plain if not model.section_responsible_ids(s.id))

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    title_font = Font(bold=True, size=16, color=HEADER_FILL)
    sub_font = Font(bold=True, size=12, color=HEADER_FILL)
    hint_font = Font(italic=True, size=9, color="5F6B7A")
    thin = Side(style="thin", color="D9DEE5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(vertical="top", wrap_text=True)

    def write_header(ws, row: int, headers: list[str], widths: list[int] | None = None) -> None:
        for col, name in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=name)
            cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, center, border
            if widths:
                ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    def colored(cell, fill_hex: str | None, default: str) -> None:
        cell.fill = PatternFill("solid", fgColor=_hex(fill_hex, default))
        cell.border = border

    def status_of(section):
        return model.status(section.status_id)

    def category_of(section):
        return model.category(section.category_id)

    wb = Workbook()

    # ------------------------------------------------------------------ Resumen
    ws = wb.active
    ws.title = SHEET_SUMMARY
    ws.column_dimensions["A"].width = 34
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16
    ws["A1"] = f"Reporte de secciones — {meta.header}"
    ws["A1"].font = title_font
    ws["A2"] = f"Generado el {now:%d/%m/%Y %H:%M}"
    ws["A2"].font = hint_font
    ws["A3"] = f"Archivo del proyecto: {model.path}" if model.path else "Proyecto en memoria"
    ws["A3"].font = hint_font

    row = 5
    ws.cell(row=row, column=1, value="Indicadores").font = sub_font
    row += 1
    in_progress = sum(1 for s in plain if 0 < s.progress < 100)
    cards = [
        ("Secciones", len(plain)),
        ("Cláusulas del pliego", len(clauses)),
        ("Relaciones", len(relations)),
        ("Avance promedio", round(stats.avg_progress)),
        ("Secciones al 100 %", stats.completed),
        ("Secciones en curso (1–99 %)", in_progress),
        ("Secciones sin avance (0 %)", sum(1 for s in plain if s.progress == 0)),
        ("Secciones sin responsable", stats.without_responsible),
    ]
    for label, value in cards:
        ws.cell(row=row, column=1, value=label).border = border
        c = ws.cell(row=row, column=2, value=value)
        c.border, c.alignment = border, center
        if label == "Avance promedio":
            c.number_format = '0"%"'
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Por estatus").font = sub_font
    row += 1
    write_header(ws, row, ["Estatus", "Secciones", "% del total", "Avance promedio"])
    row += 1
    groups: dict[int | None, list] = {}
    for s in plain:
        groups.setdefault(s.status_id, []).append(s)
    ordered = [st.id for st in model.statuses()] + [None]
    for sid in ordered:
        items = groups.get(sid, [])
        if not items:
            continue
        st = model.status(sid)
        c = ws.cell(row=row, column=1, value=st.name if st else "(sin estatus)")
        colored(c, st.color if st else None, palette.SURFACE_ALT)
        ws.cell(row=row, column=2, value=len(items)).border = border
        pct = ws.cell(row=row, column=3, value=round(100 * len(items) / len(plain)) if plain else 0)
        pct.number_format, pct.border = '0"%"', border
        avg = ws.cell(row=row, column=4, value=round(sum(s.progress for s in items) / len(items)))
        avg.number_format, avg.border = '0"%"', border
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Por responsable").font = sub_font
    row += 1
    write_header(ws, row, ["Responsable", "Nombre", "Secciones", "Avance promedio", "Al 100 %", "En curso"])
    row += 1
    for resp in model.responsibles():
        mine = [s for s in plain if resp.id in model.section_responsible_ids(s.id)]
        c = ws.cell(row=row, column=1, value=resp.code)
        colored(c, resp.color, palette.BORDER_STRONG)
        c.font = Font(bold=True, color=palette.contrast_text(resp.color).lstrip("#"))
        ws.cell(row=row, column=2, value=resp.name).border = border
        ws.cell(row=row, column=3, value=len(mine)).border = border
        avg = ws.cell(row=row, column=4, value=round(sum(s.progress for s in mine) / len(mine)) if mine else 0)
        avg.number_format, avg.border = '0"%"', border
        ws.cell(row=row, column=5, value=sum(1 for s in mine if s.progress >= 100)).border = border
        ws.cell(row=row, column=6, value=sum(1 for s in mine if 0 < s.progress < 100)).border = border
        row += 1
    if stats.without_responsible:
        c = ws.cell(row=row, column=1, value=NO_RESPONSIBLE)
        c.border = border
        ws.cell(row=row, column=3, value=stats.without_responsible).border = border
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Por categoría").font = sub_font
    row += 1
    write_header(ws, row, ["Categoría", "Secciones", "Avance promedio"])
    row += 1
    cat_groups: dict[int | None, list] = {}
    for s in sections:
        cat_groups.setdefault(s.category_id, []).append(s)
    for cid in [c.id for c in model.categories()] + [None]:
        items = cat_groups.get(cid, [])
        if not items:
            continue
        cat = model.category(cid)
        c = ws.cell(row=row, column=1, value=cat.name if cat else "(sin categoría)")
        colored(c, cat.fill_color if cat else None, palette.SURFACE_ALT)
        ws.cell(row=row, column=2, value=len(items)).border = border
        measurable = [s for s in items if not s.is_clause]
        avg = ws.cell(row=row, column=3,
                      value=round(sum(s.progress for s in measurable) / len(measurable)) if measurable else None)
        avg.number_format, avg.border = '0"%"', border
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="El avance promedio es la media simple del porcentaje de cada sección. "
                                      "Una sección con varios responsables cuenta para cada uno de ellos. "
                                      "Las cláusulas del pliego no tienen estatus ni avance y se excluyen de los "
                                      "promedios.").font = hint_font

    # ------------------------------------------------------------------ Secciones
    ws2 = wb.create_sheet(SHEET_SECTIONS)
    headers = ["N.º", "Número", "Descripción", "Categoría", "Estatus", "Avance (%)", "Responsables",
               "Observaciones", "Referencia a (n)", "Referenciada por (n)", "Última actualización"]
    write_header(ws2, 1, headers, [6, 14, 40, 24, 18, 12, 22, 50, 14, 16, 20])
    G = model.graph.G
    for i, s in enumerate(sections, start=1):
        r = i + 1
        st, cat = status_of(s), category_of(s)
        resp = ", ".join(x.code for x in model.section_responsibles(s.id))
        values = [i, s.code, s.title, cat.name if cat else "", st.name if st else "",
                  None if s.is_clause else s.progress, resp,
                  s.notes or "", G.out_degree(s.id) if s.id in G else 0, G.in_degree(s.id) if s.id in G else 0,
                  (s.updated_at or "").replace("T", " ")]
        for col, value in enumerate(values, start=1):
            cell = ws2.cell(row=r, column=col, value=value)
            cell.border = border
        ws2.cell(row=r, column=1).alignment = center
        fill, bord = model.section_colors(s)
        colored(ws2.cell(row=r, column=4), cat.fill_color if cat else fill, palette.SURFACE_ALT)
        colored(ws2.cell(row=r, column=5), st.color if st else None, palette.SURFACE_ALT)
        ws2.cell(row=r, column=6).number_format = '0"%"'
        ws2.cell(row=r, column=6).alignment = center
        ws2.cell(row=r, column=8).alignment = wrap
        ws2.cell(row=r, column=3).alignment = wrap
    last = len(sections) + 1
    if sections:
        ws2.conditional_formatting.add(f"F2:F{last}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                   end_value=100, color=BAR_COLOR, showValue=True))
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(last, 2)}"
    ws2.freeze_panes = "C2"

    # ------------------------------------------------------------------ Por responsable
    ws3 = wb.create_sheet(SHEET_BY_RESP)
    headers3 = ["Responsable", "Nombre", "Número", "Descripción", "Categoría", "Estatus", "Avance (%)", "Observaciones"]
    write_header(ws3, 1, headers3, [14, 30, 14, 40, 24, 18, 12, 50])
    r = 2
    for resp in model.responsibles():
        for s in plain:
            if resp.id not in model.section_responsible_ids(s.id):
                continue
            st, cat = status_of(s), category_of(s)
            row_values = [resp.code, resp.name, s.code, s.title, cat.name if cat else "", st.name if st else "",
                          s.progress, s.notes or ""]
            for col, value in enumerate(row_values, start=1):
                ws3.cell(row=r, column=col, value=value).border = border
            colored(ws3.cell(row=r, column=1), resp.color, palette.BORDER_STRONG)
            ws3.cell(row=r, column=1).font = Font(bold=True, color=palette.contrast_text(resp.color).lstrip("#"))
            colored(ws3.cell(row=r, column=5), cat.fill_color if cat else None, palette.SURFACE_ALT)
            colored(ws3.cell(row=r, column=6), st.color if st else None, palette.SURFACE_ALT)
            ws3.cell(row=r, column=7).number_format = '0"%"'
            ws3.cell(row=r, column=8).alignment = wrap
            r += 1
    for s in plain:
        if model.section_responsible_ids(s.id):
            continue
        st, cat = status_of(s), category_of(s)
        row_values = [NO_RESPONSIBLE, "", s.code, s.title, cat.name if cat else "", st.name if st else "",
                      s.progress, s.notes or ""]
        for col, value in enumerate(row_values, start=1):
            ws3.cell(row=r, column=col, value=value).border = border
        ws3.cell(row=r, column=7).number_format = '0"%"'
        r += 1
    if r > 2:
        ws3.conditional_formatting.add(f"G2:G{r - 1}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                    end_value=100, color=BAR_COLOR, showValue=True))
    ws3.auto_filter.ref = f"A1:{get_column_letter(len(headers3))}{max(r - 1, 2)}"
    ws3.freeze_panes = "A2"

    # ------------------------------------------------------------------ Relaciones
    ws4 = wb.create_sheet(SHEET_RELATIONS)
    headers4 = ["N.º", "Sección A", "Relación", "Sección B", "Categoría A", "Categoría B"]
    write_header(ws4, 1, headers4, [6, 40, 12, 40, 24, 24])
    for i, rel in enumerate(sorted(relations, key=lambda x: (model.section(denormalize(x)[0]).code_key
                                                             if model.section(denormalize(x)[0]) else "")), start=1):
        a_id, _kind, b_id = denormalize(rel)
        sa, sb = model.section(a_id), model.section(b_id)
        arrow = "→"
        ca, cb = category_of(sa) if sa else None, category_of(sb) if sb else None
        values = [i, sa.label if sa else "?", arrow, sb.label if sb else "?", ca.name if ca else "", cb.name if cb else ""]
        for col, value in enumerate(values, start=1):
            cell = ws4.cell(row=i + 1, column=col, value=value)
            cell.border = border
        ws4.cell(row=i + 1, column=1).alignment = center
        ws4.cell(row=i + 1, column=3).alignment = center
        colored(ws4.cell(row=i + 1, column=5), ca.fill_color if ca else None, palette.SURFACE_ALT)
        colored(ws4.cell(row=i + 1, column=6), cb.fill_color if cb else None, palette.SURFACE_ALT)
    ws4.auto_filter.ref = f"A1:F{max(len(relations) + 1, 2)}"
    ws4.freeze_panes = "A2"

    # ------------------------------------------------------------------ Mapa (opcional)
    if map_image is not None and Path(map_image).exists():
        try:
            from openpyxl.drawing.image import Image as XLImage

            ws5 = wb.create_sheet(SHEET_MAP)
            ws5["A1"] = f"Mapa de referencias — {meta.header} — {now:%d/%m/%Y}"
            ws5["A1"].font = sub_font
            img = XLImage(str(map_image))
            max_w = 1400
            if img.width > max_w:
                ratio = max_w / img.width
                img.width, img.height = int(img.width * ratio), int(img.height * ratio)
            ws5.add_image(img, "A3")
        except Exception as exc:  # noqa: BLE001 - la imagen es opcional; el reporte no debe fallar por ella
            ws5 = wb[SHEET_MAP] if SHEET_MAP in wb.sheetnames else wb.create_sheet(SHEET_MAP)
            ws5["A3"] = f"No se pudo insertar la imagen del mapa: {exc}"

    stats.sheets = list(wb.sheetnames)
    try:
        wb.save(path)
    except PermissionError as exc:
        raise ReportExportError(
            f"No se pudo escribir {path.name}. Si el archivo está abierto en Excel, ciérrelo e intente de nuevo.") from exc
    except OSError as exc:
        raise ReportExportError(f"No se pudo guardar el reporte: {exc}") from exc
    return stats
