"""Plantillas e intercambio de tablas del proyecto (secciones y relaciones) en Excel/CSV.

Formato de la plantilla:
- Hoja/archivo "Proyecto"   (opcional): Código | <valor>  y  Nombre | <valor>
- Hoja/archivo "Secciones":  Número | Descripción | Categoría | Color | Estatus | Avance | Responsables | Observaciones
- Hoja/archivo "Relaciones": Sección A | Relación | Sección B | Observaciones
Las secciones pueden escribirse como "03 30 00" o "03 30 00 - Concreto".
Cada dirección es una relación (flecha) independiente: si A referencia a B y B referencia a A, son dos filas.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Callable, Iterable

from models.entities import UiKind
from models.relation_normalizer import (
    DuplicateRelationError,
    SelfRelationError,
    code_key,
    denormalize,
    search_key,
    split_code_title,
)

if TYPE_CHECKING:
    from models.project_model import ProjectModel

SECTION_HEADERS = ["Número", "Descripción", "Categoría", "Color", "Estatus", "Avance", "Responsables", "Observaciones"]
RELATION_HEADERS = ["Sección A", "Relación", "Sección B", "Observaciones"]
PROJECT_FIELDS = [("Código", "code"), ("Nombre", "name")]
RESPONSIBLE_COLORS = ("#5B9BD5", "#70AD47", "#7030A0", "#BF9000", "#ED7D31", "#C00000", "#00B0F0", "#7F7F7F")
SHEET_PROJECT = "Proyecto"
SHEET_SECTIONS = "Secciones"
SHEET_RELATIONS = "Relaciones"
SHEET_HELP = "Instrucciones"

# Texto histórico de la plantilla 0.2.x ("Referencia mutua ↔"): hoy equivale a DOS filas, una por dirección.
KIND_BOTH = "both"

EXAMPLE_PROJECT = [("Código", "CC-26-01"), ("Nombre", "Proyecto de ejemplo")]
EXAMPLE_SECTIONS = [
    ("03 30 00", "Concreto", "Técnica / constructiva", "", "En elaboración", 70, "INIO, INIC", "Pendiente revisión de mezcla"),
    ("31 23 00", "Excavación", "Técnica / constructiva", "", "Aprobada", 100, "INIG", ""),
    ("01 31 19", "Conferencia inicial", "Contractual", "", "No iniciada", 0, "INI-PY", ""),
    ("01 35 29", "Requisitos de seguridad", "Auxiliar / apoyo", "#DDEBF7", "En revisión", 40, "INIO", ""),
    ("4.28.61", "PAGO FINAL", "Cláusula", "", "", "", "", "Cláusula del pliego: sin estatus ni avance"),
]
EXAMPLE_RELATIONS = [
    ("31 23 00 - Excavación", UiKind.REFERENCES.value, "03 30 00 - Concreto", ""),
    ("01 31 19", UiKind.REFERENCED_BY.value, "31 23 00", "Ver artículo 3.2"),
    ("01 35 29", UiKind.REFERENCES.value, "03 30 00", ""),
    ("03 30 00", UiKind.REFERENCES.value, "01 35 29", "Sentido contrario: son dos flechas"),
    ("01 31 19", UiKind.REFERENCES.value, "4.28.61", ""),
]
HELP_LINES = [
    "Plantilla de SpecRel para crear un mapa de referencias desde Excel.",
    "",
    "Cómo usarla: complete las hojas y luego ARRASTRE este archivo sobre la ventana de SpecRel",
    "(o use Archivo → Tablas → Importar / crear mapa desde Excel o CSV…).",
    "Si no hay ningún proyecto abierto, SpecRel crea el archivo del proyecto (.specrel) junto a este Excel",
    "y muestra el mapa. Si ya hay un proyecto abierto, puede agregar las filas a ese proyecto o crear uno nuevo.",
    "",
    "Hoja 'Proyecto' (opcional): Código y Nombre del proyecto nuevo. Si falta, se usa el nombre del archivo.",
    "",
    "Hoja 'Secciones': una fila por sección.",
    "  Número      -> obligatorio (ej. 03 30 00 o 4.28.33).",
    "  Descripción -> nombre de la sección (ej. Concreto).",
    "  Categoría   -> nombre de categoría; si no existe se crea. Vacío = clasificación del catálogo.",
    "  Color       -> opcional, color personalizado en hexadecimal (#RRGGBB).",
    "  Estatus     -> nombre del estatus (No iniciada, En elaboración, …); si no existe se crea.",
    "  Avance      -> porcentaje 0 a 100.",
    "  Responsables-> códigos separados por coma (INIO, INIG, …); los desconocidos se crean.",
    "  Observaciones -> texto libre.",
    "  Cláusulas del pliego (4.28.N o 4.28.N.M): se reconocen por su numeración (o Categoría = Cláusula).",
    "              Son nodos rosados sin estatus, avance ni responsables: esas columnas se ignoran.",
    "",
    "Hoja 'Relaciones': una fila por relación (flecha).",
    "  Sección A / Sección B -> número, o 'número - descripción'. Las secciones inexistentes se crean.",
    "  Relación -> 'Hace referencia a →' (A → B) o '← Es referenciada por' (B → A).",
    "              También se aceptan '->', '<-', 'A->B'.",
    "  Observaciones -> texto libre; se muestra al pasar el mouse sobre la flecha.",
    "  Cada dirección es una flecha independiente: si A referencia a B y B referencia a A,",
    "  escriba dos filas (A → B y B → A). Las filas repetidas se omiten y se informan al final.",
    "",
    "También puede usar archivos CSV separados (secciones.csv, relaciones.csv y proyecto.csv) con los mismos encabezados.",
]


class ProjectIOError(Exception):
    pass


@dataclass
class ImportSummary:
    sections_created: int = 0
    sections_updated: int = 0
    categories_created: int = 0
    statuses_created: int = 0
    responsibles_created: int = 0
    relations_created: int = 0
    relations_duplicated: int = 0
    errors: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)   # filas de relaciones omitidas, con el motivo

    def text(self, max_lines: int = 15) -> str:
        lines = [
            f"Secciones creadas: {self.sections_created}",
            f"Secciones actualizadas: {self.sections_updated}",
            f"Categorías creadas: {self.categories_created}",
            f"Estatus creados: {self.statuses_created}",
            f"Responsables creados: {self.responsibles_created}",
            f"Relaciones creadas: {self.relations_created}",
            f"Relaciones ya existentes u omitidas: {self.relations_duplicated}",
        ]
        for title, items in (("Filas omitidas", self.skipped), ("Filas con problemas", self.errors)):
            if items:
                lines.append("")
                lines.append(f"{title} ({len(items)}):")
                lines.extend(f"  • {e}" for e in items[:max_lines])
                if len(items) > max_lines:
                    lines.append(f"  … y {len(items) - max_lines} más")
        return "\n".join(lines)

    def full_text(self) -> str:
        """Detalle completo (para copiar al portapapeles)."""
        return self.text(max_lines=10_000)


# ============================================================================ plantilla
def write_template_xlsx(path: Path, with_examples: bool = True) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:  # pragma: no cover
        raise ProjectIOError("openpyxl no está instalado.") from exc

    wb = Workbook()
    ws_sec = wb.active
    ws_sec.title = SHEET_SECTIONS
    ws_rel = wb.create_sheet(SHEET_RELATIONS)
    ws_proj = wb.create_sheet(SHEET_PROJECT)
    ws_help = wb.create_sheet(SHEET_HELP)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    def write_headers(ws, headers: list[str], widths: list[int]) -> None:
        for col, (name, width) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(row=1, column=col, value=name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"

    write_headers(ws_sec, SECTION_HEADERS, [16, 40, 26, 10, 18, 10, 22, 40])
    write_headers(ws_rel, RELATION_HEADERS, [40, 28, 40, 40])
    for i, (label, _key) in enumerate(PROJECT_FIELDS, start=1):
        cell = ws_proj.cell(row=i, column=1, value=label)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E3ECF5")
    ws_proj.column_dimensions["A"].width = 14
    ws_proj.column_dimensions["B"].width = 50
    if with_examples:
        for row in EXAMPLE_SECTIONS:
            ws_sec.append(list(row))
        for row in EXAMPLE_RELATIONS:
            ws_rel.append(list(row))
        for i, (_label, value) in enumerate(EXAMPLE_PROJECT, start=1):
            ws_proj.cell(row=i, column=2, value=value)

    kinds = ",".join(k.value for k in UiKind)
    dv = DataValidation(type="list", formula1=f'"{kinds}"', allow_blank=True, showDropDown=False)
    dv.error = "Elija un tipo de relación de la lista."
    dv.prompt = "Tipo de relación"
    ws_rel.add_data_validation(dv)
    dv.add("B2:B1000")
    dv_pct = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=True)
    dv_pct.error = "El avance debe estar entre 0 y 100."
    ws_sec.add_data_validation(dv_pct)
    dv_pct.add("F2:F1000")

    for i, line in enumerate(HELP_LINES, start=1):
        ws_help.cell(row=i, column=1, value=line)
    ws_help.column_dimensions["A"].width = 110
    try:
        wb.save(path)
    except OSError as exc:
        raise ProjectIOError(f"No se pudo guardar la plantilla: {exc}") from exc


def write_template_csv(folder: Path) -> tuple[Path, Path]:
    folder.mkdir(parents=True, exist_ok=True)
    sec_path = folder / "secciones.csv"
    rel_path = folder / "relaciones.csv"
    proj_path = folder / "proyecto.csv"
    try:
        with open(sec_path, "w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh, delimiter=";")
            writer.writerow(SECTION_HEADERS)
            writer.writerows(EXAMPLE_SECTIONS)
        with open(rel_path, "w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh, delimiter=";")
            writer.writerow(RELATION_HEADERS)
            writer.writerows(EXAMPLE_RELATIONS)
        with open(proj_path, "w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh, delimiter=";")
            writer.writerows(EXAMPLE_PROJECT)
    except OSError as exc:
        raise ProjectIOError(f"No se pudo guardar la plantilla CSV: {exc}") from exc
    return sec_path, rel_path


# ============================================================================ lectura
@dataclass
class Tables:
    sections: list[dict[str, str]] = field(default_factory=list)
    relations: list[dict[str, str]] = field(default_factory=list)
    project: dict[str, str] = field(default_factory=dict)   # "code" / "name" de la hoja Proyecto

    @property
    def is_empty(self) -> bool:
        return not self.sections and not self.relations


def _norm_header(h: str) -> str:
    return search_key(str(h or "")).strip()


_SECTION_KEYS = {
    "code": ("numero", "número", "codigo", "seccion", "section", "code", "number"),
    "title": ("descripcion", "nombre", "titulo", "title", "description", "name"),
    "category": ("categoria", "category", "tipo", "clasificacion"),
    "color": ("color", "colour", "fill"),
    "status": ("estatus", "estado", "status"),
    "progress": ("avance", "progreso", "progress", "%", "porcentaje"),
    "responsibles": ("responsables", "responsable", "responsibles", "responsible", "unidad"),
    "observations": ("observaciones", "observacion", "notas", "notes", "comentarios", "observations"),
}
_RELATION_KEYS = {
    "a": ("seccion a", "a", "origen", "source", "from", "section a"),
    "kind": ("relacion", "tipo de relacion", "tipo", "relation", "kind", "conexion"),
    "b": ("seccion b", "b", "destino", "target", "to", "section b"),
    "notes": ("observaciones", "observacion", "notas", "notes", "comentarios", "observations", "comentario"),
}
_PROJECT_KEYS = {
    "code": ("codigo", "código", "code", "codigo del proyecto"),
    "name": ("nombre", "name", "proyecto", "nombre del proyecto"),
}


def _map_columns(headers: list[str], spec: dict[str, tuple[str, ...]]) -> dict[str, int] | None:
    normalized = [_norm_header(h) for h in headers]
    mapping: dict[str, int] = {}
    for field_name, candidates in spec.items():
        for idx, h in enumerate(normalized):
            if h in candidates or any(h.startswith(c) for c in candidates if len(c) > 3):
                mapping[field_name] = idx
                break
    return mapping


def _classify(headers: list[str]) -> str | None:
    rel = _map_columns(headers, _RELATION_KEYS)
    sec = _map_columns(headers, _SECTION_KEYS)
    if rel and "a" in rel and "b" in rel:
        return "relations"
    if sec and "code" in sec:
        return "sections"
    return None


def _rows_to_dicts(headers: list[str], rows: Iterable[list], spec: dict[str, tuple[str, ...]]) -> list[dict[str, str]]:
    mapping = _map_columns(headers, spec)
    out: list[dict[str, str]] = []
    for row in rows:
        values = ["" if v is None else str(v).strip() for v in row]
        if not any(values):
            continue
        record = {k: (values[i] if i < len(values) else "") for k, i in mapping.items()}
        out.append(record)
    return out


def _project_from_rows(rows: Iterable[list]) -> dict[str, str]:
    """Hoja/CSV 'Proyecto': pares clave | valor (Código, Nombre)."""
    out: dict[str, str] = {}
    for row in rows:
        values = ["" if v is None else str(v).strip() for v in row]
        if len(values) < 2 or not values[0]:
            continue
        key = _norm_header(values[0])
        for field_name, candidates in _PROJECT_KEYS.items():
            if key in candidates and values[1]:
                out[field_name] = values[1]
    return out


def _is_project_title(title: str) -> bool:
    return _norm_header(title).startswith("proyecto") or _norm_header(title) == "project"


def _read_xlsx_tables(path: Path) -> Tables:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ProjectIOError("openpyxl no está instalado.") from exc
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ProjectIOError(f"No se pudo abrir el archivo Excel: {exc}") from exc
    tables = Tables()
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            title = _norm_header(ws.title)
            if _is_project_title(ws.title):
                tables.project.update(_project_from_rows(rows))
                continue
            headers = ["" if h is None else str(h) for h in rows[0]]
            kind = _classify(headers)
            if kind is None:
                continue
            if kind == "relations" or "relac" in title:
                tables.relations.extend(_rows_to_dicts(headers, rows[1:], _RELATION_KEYS))
            else:
                tables.sections.extend(_rows_to_dicts(headers, rows[1:], _SECTION_KEYS))
    finally:
        wb.close()
    if tables.is_empty:
        raise ProjectIOError("El archivo no contiene hojas con los encabezados esperados "
                             "(Número/Descripción o Sección A/Relación/Sección B).")
    return tables


def _read_csv_rows(path: Path) -> tuple[list[str], list[list[str]]]:
    try:
        with open(path, "r", encoding="utf-8-sig", newline="") as fh:
            sample = fh.read(4096)
            fh.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(fh, dialect)
            headers = [h.strip() for h in next(reader, [])]
            rows = [r for r in reader if any(c.strip() for c in r)]
    except (OSError, UnicodeDecodeError) as exc:
        raise ProjectIOError(f"No se pudo leer el CSV: {exc}") from exc
    return headers, rows


def _read_csv_tables(paths: list[Path]) -> Tables:
    tables = Tables()
    for path in paths:
        headers, rows = _read_csv_rows(path)
        if _is_project_title(path.stem):
            tables.project.update(_project_from_rows([headers, *rows]))
            continue
        kind = _classify(headers)
        if kind == "relations":
            tables.relations.extend(_rows_to_dicts(headers, rows, _RELATION_KEYS))
        elif kind == "sections":
            tables.sections.extend(_rows_to_dicts(headers, rows, _SECTION_KEYS))
        else:
            raise ProjectIOError(f"{path.name}: encabezados no reconocidos. Use la plantilla.")
    return tables


def read_tables(paths: list[Path]) -> Tables:
    """Lee uno o varios archivos (XLSX con hojas, o CSV de secciones y/o relaciones y/o proyecto)."""
    if not paths:
        raise ProjectIOError("No se indicó ningún archivo.")
    xlsx = [p for p in paths if p.suffix.lower() in (".xlsx", ".xlsm")]
    csvs = [p for p in paths if p.suffix.lower() == ".csv"]
    tables = Tables()
    for p in xlsx:
        t = _read_xlsx_tables(p)
        tables.sections.extend(t.sections)
        tables.relations.extend(t.relations)
        tables.project.update(t.project)
    if csvs:
        t = _read_csv_tables(csvs)
        tables.sections.extend(t.sections)
        tables.relations.extend(t.relations)
        tables.project.update(t.project)
    others = [p for p in paths if p not in xlsx and p not in csvs]
    if others:
        raise ProjectIOError(f"Formato no soportado: {others[0].suffix}")
    if tables.is_empty and csvs and not xlsx:
        raise ProjectIOError("Los CSV no contienen secciones ni relaciones (solo datos del proyecto).")
    return tables


# ============================================================================ interpretación
def parse_kind(text: str) -> UiKind | str:
    """Tipo de relación de una celda. Devuelve `KIND_BOTH` para los textos antiguos de «mutua» (= dos filas)."""
    t = search_key(text or "")
    compact = t.replace(" ", "")
    if "mutua" in t or "↔" in text or "<->" in compact or "<>" in compact or "ambas" in t:
        return KIND_BOTH
    if "referenciada" in t or "←" in text or "<-" in compact or "b->a" in compact or t.startswith("es "):
        return UiKind.REFERENCED_BY
    return UiKind.REFERENCES


def _valid_hex(value: str) -> str | None:
    v = (value or "").strip().upper()
    if not v:
        return None
    if not v.startswith("#"):
        v = "#" + v
    if len(v) == 7 and all(c in "0123456789ABCDEF" for c in v[1:]):
        return v
    return None


def KEEP_STATUS(model: "ProjectModel", section) -> int | None:  # noqa: N802
    """Estatus por defecto ya asignado al crear la sección."""
    return section.status_id


ProgressFn = Callable[[int, int], None]  # (filas procesadas, total) -> None


def apply_tables(model: "ProjectModel", tables: Tables, progress: ProgressFn | None = None) -> ImportSummary:
    """Crea/actualiza secciones y relaciones en el proyecto abierto. Nunca borra nada.

    Todo se escribe en una sola transacción (`model.bulk()`); `progress` se invoca cada pocas filas
    para que la interfaz pueda mostrar el avance.
    """
    bulk = getattr(model, "bulk", None)
    if bulk is None:
        return _apply_tables(model, tables, progress)
    with bulk():
        return _apply_tables(model, tables, progress)


def _apply_tables(model: "ProjectModel", tables: Tables, on_progress: ProgressFn | None = None) -> ImportSummary:
    summary = ImportSummary()
    total = len(tables.sections) + len(tables.relations)
    done = 0

    def tick() -> None:
        nonlocal done
        done += 1
        if on_progress is not None and (done % 10 == 0 or done == total):
            on_progress(done, total)

    def resolve_category(name: str, code: str = "") -> int | None:
        name = (name or "").strip()
        if not name:
            # Sin categoría en la fila: clasificación por defecto del catálogo, o la categoría por defecto.
            from_catalog = model.catalog_category_id(code) if code else None
            if from_catalog is not None:
                return from_catalog
            default = model.default_category()
            return default.id if default else None
        before = len(model.categories())
        cat = model.category_by_name(name, create=True)
        if len(model.categories()) > before:
            summary.categories_created += 1
        return cat.id if cat else None

    def resolve_status(name: str) -> int | None:
        name = (name or "").strip()
        if not name:
            return None
        for st in model.statuses():
            if st.name.casefold() == name.casefold():
                return st.id
        st = model.add_status(name, "#D9DEE5")
        summary.statuses_created += 1
        return st.id

    def resolve_responsibles(text: str) -> list[int]:
        ids: list[int] = []
        for raw in re.split(r"[,;/|]+", text or ""):
            code = raw.strip()
            if not code:
                continue
            resp = model.responsible_by_code(code)
            if resp is None:
                color = RESPONSIBLE_COLORS[len(model.responsibles()) % len(RESPONSIBLE_COLORS)]
                resp = model.add_responsible(code.upper(), "", color)
                summary.responsibles_created += 1
            if resp.id not in ids:
                ids.append(resp.id)
        return ids

    def parse_progress(raw: str) -> int | None:
        text = (raw or "").strip().replace("%", "").replace(",", ".")
        if not text:
            return None
        try:
            value = float(text)
        except ValueError:
            return None
        if 0 < value <= 1 and "." in text:
            value *= 100  # 0.7 -> 70 %
        return max(0, min(100, int(round(value))))

    for i, row in enumerate(tables.sections, start=2):
        tick()
        raw_code = row.get("code", "")
        code, inline_title = split_code_title(raw_code)
        title = row.get("title", "").strip() or inline_title
        if not code_key(code):
            summary.errors.append(f"Secciones fila {i}: número vacío o inválido ({raw_code!r}).")
            continue
        category_id = None if model.is_clause_code(code) else resolve_category(row.get("category", ""), code)
        color = _valid_hex(row.get("color", ""))
        if row.get("color", "").strip() and color is None:
            summary.errors.append(f"Secciones fila {i}: color inválido {row.get('color')!r}, se ignora.")
        clause_row = model.is_clause_code(code) or search_key(row.get("category", "")).strip() in ("clausula", "clausulas")
        status_id = None if clause_row else resolve_status(row.get("status", ""))
        progress = parse_progress(row.get("progress", ""))
        if row.get("progress", "").strip() and progress is None:
            summary.errors.append(f"Secciones fila {i}: avance inválido {row.get('progress')!r}, se ignora.")
        responsible_ids = (resolve_responsibles(row.get("responsibles", ""))
                           if row.get("responsibles", "").strip() and not clause_row else None)
        observations = row.get("observations", "").strip() or None
        existing = model.section_by_code(code)
        is_clause = (existing.is_clause if existing is not None else False) or model.is_clause_code(code) \
            or search_key(row.get("category", "")).strip() in ("clausula", "clausulas")
        if is_clause:
            ignored = [name for name, field_name in (("Estatus", "status"), ("Avance", "progress"),
                                                     ("Responsables", "responsibles"))
                       if row.get(field_name, "").strip()]
            if ignored:
                summary.errors.append(f"Secciones fila {i}: {code} es una cláusula; se ignoran "
                                      f"{', '.join(ignored)}.")
            if existing is None:
                entry = model.catalog_entry(code)
                if entry is not None:
                    code = entry.code
                    title = title or entry.title
                section = model.add_section(code, title, kind="clause")
                if color or observations:
                    model.update_section(section.id, section.code, title, section.category_id, observations,
                                         color if color else None, None)
                summary.sections_created += 1
            else:
                new_title = title or existing.title
                new_notes = observations if observations is not None else existing.notes
                if new_title != existing.title or new_notes != existing.notes or (color and color != existing.fill_color):
                    model.update_section(existing.id, existing.code, new_title, existing.category_id, new_notes,
                                         color if color else existing.fill_color,
                                         None if color else existing.border_color)
                    summary.sections_updated += 1
            continue
        if existing is None:
            section = model.add_section(code, title, category_id)
            model.update_section(section.id, section.code, title, category_id, observations,
                                 color if color else None, None,
                                 status_id if status_id is not None else KEEP_STATUS(model, section),
                                 progress if progress is not None else 0)
            if responsible_ids:
                model.set_section_responsibles(section.id, responsible_ids)
            summary.sections_created += 1
        else:
            new_title = title or existing.title
            new_cat = category_id if row.get("category", "").strip() else existing.category_id
            new_status = status_id if status_id is not None else existing.status_id
            new_progress = progress if progress is not None else existing.progress
            new_notes = observations if observations is not None else existing.notes
            changed = (new_title != existing.title or new_cat != existing.category_id
                       or (color and color != existing.fill_color) or new_status != existing.status_id
                       or new_progress != existing.progress or new_notes != existing.notes)
            if changed:
                model.update_section(existing.id, existing.code, new_title, new_cat, new_notes,
                                     color if color else existing.fill_color, None if color else existing.border_color,
                                     new_status, new_progress)
                summary.sections_updated += 1
            if responsible_ids is not None and responsible_ids != model.section_responsible_ids(existing.id):
                model.set_section_responsibles(existing.id, responsible_ids)
                if not changed:
                    summary.sections_updated += 1

    # Para explicar cada fila omitida: (origen, destino) -> fila del archivo que creó esa flecha.
    seen_rows: dict[tuple[int, int], int] = {}

    def describe(source_id: int, target_id: int) -> str:
        sa, sb = model.section(source_id), model.section(target_id)
        return f"{sa.code if sa else '?'} → {sb.code if sb else '?'}"

    def add_directed(row_no: int, a_id: int, kind: UiKind, b_id: int, notes: str | None) -> None:
        source_id, target_id = (a_id, b_id) if kind is UiKind.REFERENCES else (b_id, a_id)
        try:
            model.add_relation(a_id, kind, b_id, notes)
            summary.relations_created += 1
            seen_rows[(source_id, target_id)] = row_no
        except DuplicateRelationError as exc:
            summary.relations_duplicated += 1
            first = seen_rows.get((source_id, target_id))
            where = f"repite la fila {first}" if first else "ya existía en el proyecto"
            summary.skipped.append(f"Relaciones fila {row_no}: {describe(source_id, target_id)} omitida ({where}).")
            if notes and not (exc.existing.notes or "").strip():
                model.set_relation_notes(exc.existing.id, notes)
        except SelfRelationError:
            summary.errors.append(f"Relaciones fila {row_no}: una sección no puede relacionarse consigo misma.")

    for i, row in enumerate(tables.relations, start=2):
        tick()
        a_text, b_text = row.get("a", ""), row.get("b", "")
        if not a_text.strip() or not b_text.strip():
            summary.errors.append(f"Relaciones fila {i}: falta Sección A o Sección B.")
            continue
        try:
            a, created_a = model.get_or_create_section(a_text)
            b, created_b = model.get_or_create_section(b_text, near=[a.id])
        except ValueError as exc:
            summary.errors.append(f"Relaciones fila {i}: {exc}")
            continue
        summary.sections_created += int(created_a) + int(created_b)
        notes = row.get("notes", "").strip() or None
        kind = parse_kind(row.get("kind", ""))
        if kind == KIND_BOTH:
            # Texto antiguo «Referencia mutua»: hoy son dos flechas independientes.
            add_directed(i, a.id, UiKind.REFERENCES, b.id, notes)
            add_directed(i, b.id, UiKind.REFERENCES, a.id, notes)
        else:
            add_directed(i, a.id, kind, b.id, notes)
    return summary


# ============================================================================ exportación
def export_tables_xlsx(model: "ProjectModel", path: Path) -> None:
    """Escribe el proyecto actual en el formato de la plantilla (editable y re-importable)."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ProjectIOError("openpyxl no está instalado.") from exc
    write_template_xlsx(path, with_examples=False)
    wb = load_workbook(path)
    ws_sec, ws_rel, ws_proj = wb[SHEET_SECTIONS], wb[SHEET_RELATIONS], wb[SHEET_PROJECT]
    meta = model.meta()
    ws_proj.cell(row=1, column=2, value=meta.code)
    ws_proj.cell(row=2, column=2, value=meta.name)
    for s in model.sections():
        cat = model.category(s.category_id)
        st = model.status(s.status_id)
        resp = ", ".join(r.code for r in model.section_responsibles(s.id))
        ws_sec.append([s.code, s.title, cat.name if cat else "", s.fill_color or "",
                       st.name if st else "", "" if s.is_clause else s.progress, resp, s.notes or ""])
    for rel in model.relations():
        a_id, kind, b_id = denormalize(rel)
        sa, sb = model.section(a_id), model.section(b_id)
        if sa is None or sb is None:
            continue
        ws_rel.append([sa.label, kind.value, sb.label, rel.notes or ""])
    try:
        wb.save(path)
    except OSError as exc:
        raise ProjectIOError(f"No se pudo guardar el archivo: {exc}") from exc
