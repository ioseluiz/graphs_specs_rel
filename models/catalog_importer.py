"""Importación de catálogos de secciones (CSV, XLSX, SQLite MasterFormat)."""
from __future__ import annotations

import csv
import sqlite3
from dataclasses import dataclass
from pathlib import Path

from models.entities import CatalogEntry
from models.relation_normalizer import code_key, search_key

CODE_HEADERS = ("codigo_display", "codigo", "code", "numero", "número", "seccion", "sección", "section")
TITLE_HEADERS = ("nombre", "titulo", "título", "title", "descripcion", "descripción", "description")
CATEGORY_HEADERS = ("categoria", "categoría", "category", "tipo", "clasificacion", "clasificación")


class CatalogImportError(Exception):
    pass


@dataclass
class ColumnMapping:
    code: str
    title: str
    category: str | None = None


@dataclass
class TablePreview:
    headers: list[str]
    rows: list[list[str]]
    suggested: ColumnMapping | None


def _guess_column(headers: list[str], candidates: tuple[str, ...]) -> str | None:
    lowered = {search_key(h): h for h in headers}
    for cand in candidates:
        if cand in lowered:
            return lowered[cand]
    for cand in candidates:
        for key, original in lowered.items():
            if cand in key:
                return original
    return None


def suggest_mapping(headers: list[str]) -> ColumnMapping | None:
    code = _guess_column(headers, CODE_HEADERS)
    title = _guess_column(headers, TITLE_HEADERS)
    if code is None and len(headers) >= 1:
        code = headers[0]
    if title is None and len(headers) >= 2:
        title = headers[1]
    if code is None or title is None:
        return None
    return ColumnMapping(code, title, _guess_column(headers, CATEGORY_HEADERS))


def _rows_to_entries(headers: list[str], rows: list[list[str]], mapping: ColumnMapping) -> list[CatalogEntry]:
    idx = {h: i for i, h in enumerate(headers)}
    if mapping.code not in idx or mapping.title not in idx:
        raise CatalogImportError("Las columnas seleccionadas no existen en el archivo.")
    ci, ti = idx[mapping.code], idx[mapping.title]
    gi = idx.get(mapping.category) if mapping.category else None
    seen: dict[str, CatalogEntry] = {}
    for row in rows:
        if ci >= len(row):
            continue
        code = " ".join(str(row[ci]).split())
        key = code_key(code)
        if not key:
            continue
        title = str(row[ti]).strip() if ti < len(row) else ""
        category = str(row[gi]).strip() if gi is not None and gi < len(row) and row[gi] else None
        seen[key] = CatalogEntry(key, code, title, category or None)
    return list(seen.values())


# --------------------------------------------------------------------------- CSV
def preview_csv(path: Path, limit: int = 20) -> TablePreview:
    headers, rows = _read_csv(path, limit)
    return TablePreview(headers, rows, suggest_mapping(headers))


def _read_csv(path: Path, limit: int | None) -> tuple[list[str], list[list[str]]]:
    try:
        with open(path, "r", encoding="utf-8-sig", newline="") as fh:
            sample = fh.read(4096)
            fh.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(fh, dialect)
            headers = [h.strip() for h in next(reader, [])]
            rows: list[list[str]] = []
            for row in reader:
                if not any(cell.strip() for cell in row):
                    continue
                rows.append(row)
                if limit is not None and len(rows) >= limit:
                    break
    except (OSError, UnicodeDecodeError) as exc:
        raise CatalogImportError(f"No se pudo leer el CSV: {exc}") from exc
    if not headers:
        raise CatalogImportError("El CSV no tiene encabezados.")
    return headers, rows


def import_csv(path: Path, mapping: ColumnMapping) -> list[CatalogEntry]:
    headers, rows = _read_csv(path, None)
    return _rows_to_entries(headers, rows, mapping)


# --------------------------------------------------------------------------- XLSX
def _read_xlsx(path: Path, limit: int | None) -> tuple[list[str], list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise CatalogImportError("openpyxl no está instalado.") from exc
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise CatalogImportError(f"No se pudo abrir el XLSX: {exc}") from exc
    ws = wb.active
    headers: list[str] = []
    rows: list[list[str]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        values = ["" if v is None else str(v) for v in row]
        if i == 0:
            headers = [h.strip() for h in values]
            continue
        if not any(v.strip() for v in values):
            continue
        rows.append(values)
        if limit is not None and len(rows) >= limit:
            break
    wb.close()
    if not headers:
        raise CatalogImportError("La hoja no tiene encabezados.")
    return headers, rows


def preview_xlsx(path: Path, limit: int = 20) -> TablePreview:
    headers, rows = _read_xlsx(path, limit)
    return TablePreview(headers, rows, suggest_mapping(headers))


def import_xlsx(path: Path, mapping: ColumnMapping) -> list[CatalogEntry]:
    headers, rows = _read_xlsx(path, None)
    return _rows_to_entries(headers, rows, mapping)


# --------------------------------------------------------------------------- SQLite MasterFormat
def import_masterformat_db(path: Path, max_level: int | None = 3) -> list[CatalogEntry]:
    """Lee la tabla `master_format` (codigo_display, nombre, nivel) de 02_master_format_database."""
    try:
        conn = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True)
    except sqlite3.Error as exc:
        raise CatalogImportError(f"No se pudo abrir la base: {exc}") from exc
    try:
        cols = {r[1] for r in conn.execute("PRAGMA table_info(master_format)")}
        if not {"codigo_display", "nombre"} <= cols:
            raise CatalogImportError("La base no contiene la tabla master_format esperada.")
        sql = "SELECT codigo_display, nombre FROM master_format"
        params: tuple = ()
        if max_level is not None and "nivel" in cols:
            sql += " WHERE nivel <= ?"
            params = (max_level,)
        entries: dict[str, CatalogEntry] = {}
        for code, title in conn.execute(sql, params):
            code = " ".join(str(code or "").split())
            key = code_key(code)
            if key:
                entries[key] = CatalogEntry(key, code, str(title or "").strip(), None)
        return list(entries.values())
    except sqlite3.Error as exc:
        raise CatalogImportError(str(exc)) from exc
    finally:
        conn.close()


def preview_for(path: Path) -> TablePreview:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return preview_csv(path)
    if suffix in (".xlsx", ".xlsm"):
        return preview_xlsx(path)
    raise CatalogImportError(f"Formato no soportado para vista previa: {suffix}")


def import_for(path: Path, mapping: ColumnMapping | None, max_level: int | None = 3) -> list[CatalogEntry]:
    suffix = path.suffix.lower()
    if suffix in (".db", ".sqlite", ".sqlite3"):
        return import_masterformat_db(path, max_level)
    if mapping is None:
        raise CatalogImportError("Debe indicar las columnas de código y título.")
    if suffix == ".csv":
        return import_csv(path, mapping)
    if suffix in (".xlsx", ".xlsm"):
        return import_xlsx(path, mapping)
    raise CatalogImportError(f"Formato no soportado: {suffix}")
